# Copyright (c) 2026, IT Team - TIQN and contributors
# For license information, please see license.txt

"""Desk side of payslips: Excel template, upload (preview -> import as Draft),
publish and withdraw. Used by the `payslip-upload` page (the TIQN Payslip list has no
import / publish buttons since 27/10/2026). CnB (and System Manager) only.

Excel layout (one row per employee; template since 24/10/2026):

    No. | Period | Period From | Period To | Pay Date | Employee | Employee Name |
    Designation | Department | Section | Group | Note | <components>

No. and Designation are for HR's eyes only - never read back. Note (27/10/2026) = optional
free text for the whole payslip (Mini App + printout), at most NOTE_MAX characters. The older Vietnamese
headers (Kỳ lương | MSNV | Họ tên | dept | custom_section | custom_group) are still read.
Period From / To may be cut inside the period (joined / left during it); empty = the
controller works them out. Pay Date empty = the monthly pay day (or relieving + 14 days).

Two templates (download_template):
- monthly: everyone on the payroll for the WHOLE end of the period - not those who left
  before or during it (left during it = final settlement, paid separately), not those on
  maternity leave for the whole period;
- resigned: everyone whose relieving_date falls in a date range HR picks, each in the
  period of their last working day, Period To and Pay Date filled in.

Line columns follow the TIQN Payslip Structure of the period (28/10/2026, see
payslip_structure.py) and are matched by the leading code of the header
("3.2 Technical allowance/ Phụ cấp kỹ thuật" -> line "3.2"), so HR may keep the long
bilingual names, reorder columns or fix a typo in the text part. "<code> Days" /
"<code> Hours" (older files: "... - Days" / "Unit/ĐVT") is the Unit/ĐVT number of that line
- e.g. the 27 days of 1.2; the template puts it just before the amount. All periods of one
file must use the same structure. A line missing from the file is 0 (listed for HR).

Numbers are taken exactly as HR computed them - no formula is checked (user, 28/10/2026).
Only impossible numbers are refused (payslip_structure.absurd_errors); a negative net
salary is a warning HR confirms.
"""

import io
import re
import unicodedata

import frappe
from frappe import _
from frappe.utils import add_days, cint, flt, format_datetime, formatdate, getdate, now_datetime, today

from customize_erpnext.customize_erpnext.doctype.attendance_calculation_setting.attendance_calculation_setting import (
	get_excluded_employee_ids,
)
from customize_erpnext.employee_portal import payslip_structure as pss
from customize_erpnext.employee_portal.doctype.tiqn_payslip.tiqn_payslip import FINAL_SETTLEMENT_PAY_DAYS
from customize_erpnext.employee_portal.notify import notify_all
from customize_erpnext.employee_portal.utils import (
	parse_period_cell,
	period_bounds,
	period_label,
	period_of_date,
	shift_period,
	strip_company_suffix,
	validate_period,
)

ROLES = ("CnB", "System Manager")
QUEUE_THRESHOLD = 150

IDENT_COLUMNS = {
	"period": ("ky luong", "ky", "period", "thang"),
	"period_from": ("period from", "tu ngay"),
	"period_to": ("period to", "den ngay"),
	"pay_date": ("pay date", "ngay tra luong", "ngay thanh toan"),
	"employee": ("msnv", "ma nv", "ma nhan vien", "employee", "employee id"),
	"employee_name": ("ho ten", "ho va ten", "ten", "employee name", "full name"),
	"department": ("dept", "department", "phong ban"),
	"section": ("custom_section", "section"),
	"group": ("custom_group", "group", "to"),
	"note": ("note", "notes", "ghi chu", "remark", "remarks"),
}
IGNORED_COLUMNS = ("stt", "no", "no.", "#", "designation", "chuc danh", "chuc vu")
TEMPLATE_IDENT = (
	"No.", "Period", "Period From", "Period To", "Pay Date", "Employee", "Employee Name",
	"Designation", "Department", "Section", "Group", "Note",
)
NOTE_MAX = 1000
CODE_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)(?=\s|$|/|\(|-)")
QTY_LABELS = {"Day": "Days", "Hour": "Hours"}
# On the _norm()-ed header, when it is not exactly a field label: "<code> Day(s)/Hour(s)"
# alone, or "- Days" at the END (older files) - never a bare "days" inside a name:
# "1.1 Standard Working Days" is not a days column.
QTY_RE = re.compile(r"^\d+(?:\.\d+)*\s+(?:days?|hours?)$|-\s*(?:days?|hours?)\s*$|unit\s*/\s*dvt|\bdvt\b")


def amount_header(line):
	"""Template header of a line: "<code> <English label>"."""
	return f"{line.code} {line.label_en}".strip()


def qty_header(line):
	"""Template header of the Unit/ĐVT number of a line: "1.2 Days" / "2.1 Hours"."""
	return f"{line.code} {QTY_LABELS.get(line.quantity_unit, 'Days')}"


def _check_role():
	frappe.only_for(ROLES)


def _norm(text):
	text = unicodedata.normalize("NFD", str(text or "")).replace("đ", "d").replace("Đ", "D")
	text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
	return re.sub(r"\s+", " ", text).strip().lower()


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------
@frappe.whitelist(methods=["GET"])
def get_template_periods():
	"""Periods offered by the "Download Template" dialog: the current one and the 12
	before it. Default = the last period that has ENDED (the one HR is computing)."""
	_check_role()
	current = period_of_date(today())
	out = []
	for i in range(13):
		p = shift_period(current, -i)
		start, end = period_bounds(p)
		out.append(
			{
				"value": p,
				"label": f"{period_label(p)} ({formatdate(start)} - {formatdate(end)})",
				"from_date": str(start),
				"to_date": str(end),
			}
		)
	return {"default": shift_period(current, -1), "periods": out}


@frappe.whitelist(methods=["GET"])
def get_report_periods():
	"""Pay Period filter of the payslip reports: the periods that have payslips, newest
	first. Default = the newest one."""
	_check_role()
	periods = sorted(frappe.get_all("TIQN Payslip", pluck="period", distinct=True), reverse=True)
	out = []
	for p in periods:
		start, end = period_bounds(p)
		out.append({"value": p, "label": f"{period_label(p)} ({formatdate(start)} - {formatdate(end)})"})
	return {"default": periods[0] if periods else None, "periods": out}


def template_employees(period):
	"""Monthly template: on the payroll until the END of the period. Not those who left
	before or during it (🔴 relieving_date = first day OFF; left during it = final
	settlement, see resigned_employees), not those on maternity leave for the whole
	period (no working day in it), not Attendance Calculation Setting > Exclude Employees."""
	start, end = period_bounds(period)
	return _not_excluded(frappe.db.sql(
		"""
		SELECT e.name, e.employee_name, e.designation, e.department, e.custom_section, e.custom_group,
		       e.date_of_joining, e.relieving_date
		FROM `tabEmployee` e
		WHERE (e.date_of_joining IS NULL OR e.date_of_joining <= %(end)s)
		  AND (e.relieving_date IS NULL OR e.relieving_date > %(end)s)
		  AND NOT EXISTS (
			SELECT 1 FROM `tabEmployee Maternity` m
			WHERE m.employee = e.name
			  AND m.maternity_from_date <= %(start)s AND m.maternity_to_date >= %(end)s
		  )
		ORDER BY e.name
		""",
		{"start": start, "end": end},
		as_dict=True,
	))


def resigned_employees(from_date, to_date):
	"""Resigned template: relieving_date in [from_date, to_date], whatever the period.
	Attendance Calculation Setting > Exclude Employees left out too, and so are those
	whose last working day is the LAST day of a period (relieving_date = the 26th): they
	worked the whole period, so they are on the monthly template (paid on the 5th, still
	within the 14 days) - listing them here too would make a second payslip (26/10/2026)."""
	return [e for e in _resigned_in_range(from_date, to_date) if not _left_at_period_end(e.relieving_date)]


def _left_at_period_end(relieving_date):
	last = add_days(relieving_date, -1)
	return getdate(last) == period_bounds(period_of_date(last))[1]


def _resigned_in_range(from_date, to_date):
	return _not_excluded(frappe.db.sql(
		"""
		SELECT name, employee_name, designation, department, custom_section, custom_group,
		       date_of_joining, relieving_date
		FROM `tabEmployee`
		WHERE relieving_date BETWEEN %(from_date)s AND %(to_date)s
		  AND (date_of_joining IS NULL OR date_of_joining < relieving_date)
		ORDER BY relieving_date, name
		""",
		{"from_date": getdate(from_date), "to_date": getdate(to_date)},
		as_dict=True,
	))


def _not_excluded(employees):
	"""Drop Attendance Calculation Setting > Exclude Employees (test codes, people not on
	this payroll) - the same list headcount and attendance leave out."""
	excluded = get_excluded_employee_ids()
	return [e for e in employees if e.name not in excluded]


def template_row(e, period=None):
	"""Period, From, To, Pay Date of one template row. Resigned (no `period`): the
	period of the last working day, cut at relieving_date - 1, paid 14 days after."""
	if period:
		start, end = period_bounds(period)
		last, pay = end, None
	else:
		last = add_days(e.relieving_date, -1)
		period = period_of_date(last)
		start, end = period_bounds(period)
		pay = add_days(e.relieving_date, FINAL_SETTLEMENT_PAY_DAYS)
	first = max(start, getdate(e.date_of_joining)) if e.date_of_joining else start
	return period, first, last, pay


@frappe.whitelist(methods=["GET"])
def download_template(period=None, resigned=0, from_date=None, to_date=None):
	"""Excel template, employee columns filled in - HR only types the numbers.
	`period`: monthly template. `resigned` + `period` (+ optional from/to inside it, empty =
	the whole period): employees whose relieving date falls in that range - every row then
	belongs to `period` (26/10/2026: the range can no longer span periods).
	Neither: header + one example row."""
	_check_role()
	resigned = cint(resigned)
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill

	if period or resigned:
		try:
			period = validate_period(period)
		except frappe.ValidationError:
			frappe.throw(_("Period must look like YYYY-MM."))
	if resigned:
		start, end = period_bounds(period)
		from_date = getdate(from_date) if from_date else start
		to_date = getdate(to_date) if to_date else end
		if not start <= from_date <= to_date <= end:
			frappe.throw(
				_("The relieving date range must be inside period {0} ({1} - {2}), From <= To.").format(
					period_label(period), formatdate(start), formatdate(end)
				)
			)

	if period or resigned:
		_structure, lines = pss.period_structure_lines(period)
	else:
		latest = frappe.db.get_value(pss.STRUCTURE_DT, {}, "name", order_by="valid_from desc")
		if not latest:
			frappe.throw(_("Create a Payslip Structure first."))
		lines = pss.structure_lines(latest)

	wb = Workbook()
	ws = wb.active
	ws.title = "Payslip"
	headers = list(TEMPLATE_IDENT)
	for line in lines:
		if line.quantity_unit:
			headers.append(qty_header(line))
		headers.append(amount_header(line))
	ws.append(headers)
	blanks = [None] * (len(headers) - len(TEMPLATE_IDENT))
	if period or resigned:
		employees = resigned_employees(from_date, to_date) if resigned else template_employees(period)
		for no, e in enumerate(employees, start=1):
			row_period, first, last, pay = template_row(e, None if resigned else period)
			ws.append(
				[no, row_period, first, last, pay, e.name, e.employee_name, e.designation,
				 strip_company_suffix(e.department), e.custom_section, e.custom_group, None]
				+ blanks
			)
	else:
		example_from, example_to = period_bounds("2026-10")
		ws.append([1, "2026-10", example_from, example_to, None, "TIQN-0001", "", "", "", "", "", ""] + blanks)
	bold = Font(bold=True)
	fill = PatternFill("solid", fgColor="DDEBF7")
	for cell in ws[1]:
		cell.font = bold
		cell.fill = fill
		cell.alignment = Alignment(wrap_text=True, vertical="top")
	ws.row_dimensions[1].height = 60
	for i in range(1, len(headers) + 1):
		ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 14 if i > len(TEMPLATE_IDENT) else 16
	# No. | Period | From | To | Pay Date | Employee | Name | Designation | Dept | Section | Group | Note
	for letter, width in zip("ABCDEFGHIJKL", (6, 9, 11, 11, 11, 11, 26, 10, 10, 9, 9, 10)):
		ws.column_dimensions[letter].width = width
	# Frozen up to Employee Name (27/10/2026): who the row is stays in view while scrolling.
	ws.freeze_panes = ws.cell(row=2, column=TEMPLATE_IDENT.index("Employee Name") + 2).coordinate
	# Period and Employee must stay text, or Excel turns "10-2026" into a date and
	# "0148" into 148; Period From / To and Pay Date are real dates.
	for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 1000)):
		row[1].number_format = row[5].number_format = "@"
		row[2].number_format = row[3].number_format = row[4].number_format = "dd/mm/yyyy"

	out = io.BytesIO()
	wb.save(out)
	if period or resigned:
		# Holds the employee list of that moment: data file -> time stamp (app-wide rule).
		what = f"{period} Resigned {from_date.strftime('%y%m%d')}-{to_date.strftime('%y%m%d')}" if resigned else period
		filename = f"Payslip {what} {now_datetime().strftime('%y%m%d %H%M%S')}.xlsx"
	else:
		filename = "payslip_template.xlsx"
	frappe.response["filename"] = filename
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


# ---------------------------------------------------------------------------
# Parse + validate
# ---------------------------------------------------------------------------
def _read_rows(file_url):
	from openpyxl import load_workbook

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	if not (file_doc.file_name or "").lower().endswith(".xlsx"):
		frappe.throw(_("Please upload an .xlsx file."))
	wb = load_workbook(io.BytesIO(file_doc.get_content()), read_only=True, data_only=True)
	ws = wb.worksheets[0]
	return [list(r) for r in ws.iter_rows(values_only=True)]


def _find_header(rows):
	for i, row in enumerate(rows[:20]):
		names = {_norm(c) for c in row if c not in (None, "")}
		if names & set(IDENT_COLUMNS["employee"]):
			return i
	frappe.throw(_("Could not find the header row. It must contain an MSNV column."))


def _map_columns(header, lines, structure):
	"""mapping["components"] = {column index: (line, "amount" | "quantity", is_qty)}."""
	by_code = {line.code: line for line in lines}
	mapping = {"ident": {}, "components": {}}
	errors, seen = [], {}
	for idx, raw in enumerate(header):
		if raw in (None, ""):
			continue
		norm = _norm(raw)
		ident = next((k for k, names in IDENT_COLUMNS.items() if norm in names), None)
		if ident and ident not in mapping["ident"]:
			mapping["ident"][ident] = idx
			continue
		if norm in IGNORED_COLUMNS:
			continue
		m = CODE_RE.match(str(raw))
		line = by_code.get(m.group(1)) if m else None
		if not line:
			errors.append(_("Column \"{0}\" does not match any line of Payslip Structure {1}.").format(raw, structure or "-"))
			continue
		# "1.2 Days" (Unit/ĐVT number) vs "1.1 Standard Working Days" (a days-only line,
		# despite the "Days"). Older files: "<bilingual label> - Days" or "Unit/ĐVT" (QTY_RE).
		if line.quantity_unit and norm == _norm(qty_header(line)):
			is_qty = True
		elif norm == _norm(amount_header(line)):
			is_qty = False
		else:
			is_qty = bool(QTY_RE.search(norm))
		if is_qty and not line.quantity_unit:
			errors.append(_("Column \"{0}\": line {1} has no Unit/ĐVT column.").format(raw, line.code))
			continue
		key = (line.code, is_qty)
		if key in seen:
			errors.append(_("Columns \"{0}\" and \"{1}\" are both line {2}.").format(seen[key], raw, line.code))
			continue
		seen[key] = raw
		target = "quantity" if is_qty or line.unit != "VND" else "amount"
		mapping["components"][idx] = (line, target, is_qty)
	for required in ("period", "employee"):
		if required not in mapping["ident"]:
			errors.append(_("Missing column {0}.").format({"period": "Kỳ lương", "employee": "MSNV"}[required]))
	missing = [amount_header(line) for line in lines if (line.code, False) not in seen]
	missing += [qty_header(line) for line in lines if line.quantity_unit and (line.code, True) not in seen]
	return mapping, errors, missing


def _file_structure(rows, h):
	"""(structure, lines, error) of the periods in the file: they must share one structure."""
	header = rows[h]
	col = next((i for i, raw in enumerate(header) if raw not in (None, "") and _norm(raw) in IDENT_COLUMNS["period"]), None)
	if col is None:
		return None, [], _("Missing column {0}.").format("Kỳ lương")
	periods = sorted({p for p in (parse_period_cell(r[col]) if col is not None and col < len(r) else None for r in rows[h + 1 :]) if p})
	if not periods:
		return None, [], None  # the rows report the unreadable periods
	found = {}
	for p in periods:
		found.setdefault(pss.structure_for_period(p), []).append(p)
	if None in found:
		return None, [], _("No Payslip Structure covers period(s) {0}. Create one first.").format(", ".join(found[None]))
	if len(found) > 1:
		return None, [], _("The file mixes periods of different Payslip Structures ({0}). Upload one file per structure.").format(
			"; ".join(f"{s}: {', '.join(ps)}" for s, ps in found.items())
		)
	structure = next(iter(found))
	return structure, pss.structure_lines(structure), None


def _resolve_employee(value, cache):
	text = str(value or "").strip()
	if isinstance(value, float) and value.is_integer():
		text = str(int(value))
	if not text:
		return None
	if text in cache:
		return cache[text]
	name = text if frappe.db.exists("Employee", text) else None
	if not name and text.isdigit():
		guess = f"TIQN-{int(text):04d}"
		name = guess if frappe.db.exists("Employee", guess) else None
	cache[text] = name
	return name


def _date_cell(value):
	"""A Period From / To cell -> date, None when empty, False when unreadable."""
	if value in (None, ""):
		return None
	try:
		return getdate(value)
	except Exception:
		return False


def _number(value):
	if value in (None, ""):
		return 0.0, True
	if isinstance(value, (int, float)):
		return float(value), True
	text = str(value).strip().replace(",", "")
	try:
		return float(text), True
	except ValueError:
		return None, False


def parse_upload(file_url):
	rows = _read_rows(file_url)
	if not rows:
		frappe.throw(_("The file is empty."))
	h = _find_header(rows)
	structure, lines, structure_error = _file_structure(rows, h)
	if not structure and not structure_error:
		structure_error = _("Cannot read any period in the Period column (MM-YYYY, e.g. 10-2026).")
	if structure_error:
		# Without the structure every line column would also be "unknown" - report the cause only.
		return {"structure": None, "header_errors": [structure_error], "missing_components": [], "rows": [],
			"periods": [], "error_count": 1, "warning_count": 0}
	mapping, header_errors, missing_components = _map_columns(rows[h], lines, structure)

	result = {
		"structure": structure,
		"header_errors": header_errors,
		"missing_components": missing_components,
		"rows": [],
		"periods": [],
		"error_count": 0,
		"warning_count": 0,
	}
	if header_errors:
		result["error_count"] = len(header_errors)
		return result

	ident = mapping["ident"]
	emp_cache, seen = {}, {}
	employees = {}
	existing = {}

	for offset, row in enumerate(rows[h + 1 :], start=h + 2):
		def cell(key):
			idx = ident.get(key)
			return row[idx] if idx is not None and idx < len(row) else None

		raw_emp = cell("employee")
		if raw_emp in (None, ""):
			continue
		errors, warnings = [], []
		period = parse_period_cell(cell("period"))
		if not period:
			errors.append(_("Cannot read the period \"{0}\" - use MM-YYYY, e.g. 10-2026.").format(cell("period") or ""))
		employee = _resolve_employee(raw_emp, emp_cache)
		if not employee:
			errors.append(_("Employee {0} does not exist.").format(raw_emp))
		dates = {}
		for key, label in (("period_from", "Period From"), ("period_to", "Period To"), ("pay_date", "Pay Date")):
			given = _date_cell(cell(key))
			if given is False:
				errors.append(_("{0}: cannot read the date \"{1}\".").format(_(label), cell(key)))
				given = None
			dates[key] = given
		# Free text for the whole payslip (27/10/2026). Empty / no column = no note: like the
		# amounts, nothing is kept from an earlier upload.
		note = str(cell("note") or "").strip() or None
		if note and len(note) > NOTE_MAX:
			errors.append(_("Note is longer than {0} characters.").format(NOTE_MAX))
		if period:
			start, end = period_bounds(period)
			first, last = dates["period_from"] or start, dates["period_to"] or end
			if not start <= first <= last <= end:
				errors.append(
					_("Period From / To must be inside period {0} ({1} - {2}).").format(period, formatdate(start), formatdate(end))
				)

		values = {}
		for idx, (line, target, is_qty) in mapping["components"].items():
			raw = row[idx] if idx < len(row) else None
			value, ok = _number(raw)
			if not ok:
				label = qty_header(line) if is_qty else line.code
				errors.append(_("{0}: \"{1}\" is not a number.").format(label, raw))
				continue
			values.setdefault(line.code, {})[target] = value
		snapshot = pss.make_snapshot(lines, values)
		net_pay = pss.key_values(snapshot)["net_pay"]
		if period:
			errors += pss.absurd_errors(snapshot, period)
		if net_pay < 0:
			warnings.append(_("Net salary is negative."))

		if employee and period:
			key = (employee, period)
			if key in seen:
				errors.append(_("Duplicate of row {0} (same employee and period).").format(seen[key]))
			seen[key] = offset
			if key not in existing:
				existing[key] = frappe.db.get_value(
					"TIQN Payslip",
					{"employee": employee, "period": period},
					["status", "import_batch", "modified", "period_from", "period_to", "pay_date", "note", "structure"],
					as_dict=True,
				) or frappe._dict()
			old = existing[key]
			if old.structure and structure and old.structure != structure:
				errors.append(
					_("The {0} payslip uses Payslip Structure {1}, this file {2}. Delete or withdraw it first.").format(
						period_label(period), old.structure, structure
					)
				)
			if old.status == "Published":
				errors.append(_("The {0} payslip is already published. Withdraw it first.").format(period_label(period)))
			elif old.status:
				# Draft / Withdrawn would be replaced silently - e.g. a resigned file imported
				# over the monthly one. HR confirms like any other warning (26/10/2026).
				warnings.append(
					_("Overwrites the existing {0} payslip of {1} (batch {2}, saved {3}).").format(
						_(old.status, context="Payslip"), period_label(period), old.import_batch or "-", format_datetime(old.modified)
					)
				)
				for field, label in (("period_from", "Period From"), ("period_to", "Period To"), ("pay_date", "Pay Date")):
					if dates.get(field) and old.get(field) and getdate(dates[field]) != getdate(old[field]):
						warnings.append(
							_("{0}: {1} now, {2} in the file.").format(_(label), formatdate(old[field]), formatdate(dates[field]))
						)
				if (old.note or None) != note:
					warnings.append(_("Note: \"{0}\" now, \"{1}\" in the file.").format(old.note or "-", note or "-"))

		if employee:
			if employee not in employees:
				employees[employee] = frappe.db.get_value(
					"Employee", employee, ["employee_name", "status", "relieving_date"], as_dict=True
				)
			emp = employees[employee]
			name_in_file = str(cell("employee_name") or "").strip()
			if name_in_file and _norm(name_in_file) != _norm(emp.employee_name):
				warnings.append(_("Name in file \"{0}\" differs from ERP \"{1}\".").format(name_in_file, emp.employee_name))
			if period and emp.relieving_date and getdate(emp.relieving_date) <= period_bounds(period)[0]:
				warnings.append(_("Employee left on {0}, before this period started.").format(emp.relieving_date))
			elif dates["period_to"] and emp.relieving_date and getdate(dates["period_to"]) >= getdate(emp.relieving_date):
				# Monthly file downloaded before the relieving date was entered: the controller
				# keeps the file's Period To, so it would stay past the last working day (26/10/2026).
				warnings.append(
					_("Employee left on {0} (last working day {1}), but Period To in the file is {2}.").format(
						formatdate(emp.relieving_date),
						formatdate(add_days(emp.relieving_date, -1)),
						formatdate(dates["period_to"]),
					)
				)

		result["rows"].append(
			{
				"row": offset,
				"employee": employee,
				"employee_raw": str(raw_emp),
				"employee_name": str(cell("employee_name") or "").strip() or None,
				"department": str(cell("department") or "").strip() or None,
				"section": str(cell("section") or "").strip() or None,
				"group": str(cell("group") or "").strip() or None,
				"note": note,
				"period": period,
				**dates,
				"net_pay": net_pay,
				"values": values,
				"errors": errors,
				"warnings": warnings,
				"existing_status": (existing.get((employee, period)) or {}).get("status") if employee and period else None,
			}
		)
		result["error_count"] += len(errors)
		result["warning_count"] += len(warnings)

	result["periods"] = sorted({r["period"] for r in result["rows"] if r["period"]})
	if not result["rows"]:
		result["header_errors"].append(_("No employee rows found under the header."))
		result["error_count"] += 1
	return result


@frappe.whitelist(methods=["POST"])
def preview_upload(file_url):
	_check_role()
	data = parse_upload(file_url)
	# The page shows a summary; the full component list per row is not needed there.
	for r in data["rows"]:
		r["component_count"] = len(r.pop("values"))
	return data


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------
@frappe.whitelist(methods=["POST"])
def import_upload(file_url, confirm_warnings=0):
	_check_role()
	data = parse_upload(file_url)
	if data["error_count"]:
		frappe.throw(_("The file still has {0} error(s). Fix them and upload again.").format(data["error_count"]))
	if data["warning_count"] and not cint(confirm_warnings):
		frappe.throw(_("Please confirm the warnings before importing."))

	batch = "UP-" + now_datetime().strftime("%Y%m%d-%H%M%S")
	if len(data["rows"]) > QUEUE_THRESHOLD:
		frappe.enqueue(
			_import_rows,
			queue="long",
			timeout=1800,
			rows=data["rows"],
			batch=batch,
			file_url=file_url,
			user=frappe.session.user,
			structure=data["structure"],
		)
		return {"queued": True, "batch": batch, "rows": len(data["rows"])}
	count = _import_rows(data["rows"], batch, file_url, frappe.session.user, data["structure"])
	return {"queued": False, "batch": batch, "rows": count}


def _import_rows(rows, batch, file_url, user, structure):
	count = 0
	lines = pss.structure_lines(structure)
	for r in rows:
		name = frappe.db.get_value("TIQN Payslip", {"employee": r["employee"], "period": r["period"]})
		doc = frappe.get_doc("TIQN Payslip", name) if name else frappe.new_doc("TIQN Payslip")
		if doc.status == "Published":
			# Re-checked here: it may have been published between preview and import.
			frappe.throw(_("{0} {1} is already published.").format(r["employee"], r["period"]))
		# A Withdrawn payslip was published once: employees may have opened it. Keep it
		# Withdrawn (Publish takes Withdrawn too) - the page shows it was published before.
		was_published = doc.status == "Withdrawn"
		doc.update(
			{
				"employee": r["employee"],
				"employee_name": r["employee_name"] or frappe.db.get_value("Employee", r["employee"], "employee_name"),
				"department": r["department"],
				"section": r["section"],
				"group": r["group"],
				"period": r["period"],
				# Empty = worked out again by the controller (joined / left during the period).
				"period_from": r.get("period_from"),
				"period_to": r.get("period_to"),
				"pay_date": r.get("pay_date"),
				"note": r.get("note"),
				"status": "Withdrawn" if was_published else "Draft",
				"import_batch": batch,
			}
		)
		# The structure's lines + the file's numbers; a line missing from the file is 0,
		# never left over from an earlier upload.
		doc.structure = structure
		doc.lines_json = pss.dump_snapshot(pss.make_snapshot(lines, r["values"]))
		doc.save(ignore_permissions=True) if name else doc.insert(ignore_permissions=True)
		count += 1
	_delete_upload_file(file_url)
	if frappe.flags.in_test:
		return count
	frappe.db.commit()
	frappe.publish_realtime(
		"payslip_import_done", {"batch": batch, "rows": count}, user=user, after_commit=True
	)
	return count


def _delete_upload_file(file_url):
	"""The uploaded workbook holds every salary of the period - do not leave it in File."""
	for name in frappe.get_all("File", {"file_url": file_url}, pluck="name"):
		frappe.delete_doc("File", name, ignore_permissions=True, force=True)


@frappe.whitelist(methods=["POST"])
def discard_upload(file_url):
	_check_role()
	_delete_upload_file(file_url)


# ---------------------------------------------------------------------------
# Publish / withdraw
# ---------------------------------------------------------------------------
@frappe.whitelist(methods=["GET"])
def get_period_summary():
	_check_role()
	rows = frappe.db.sql(
		"""
		SELECT period, status, COUNT(*) AS n, SUM(net_pay) AS net_total, COUNT(first_viewed_on) AS viewed,
		       SUM(final_settlement) AS final_settlement, MAX(published_on) AS published_on, MAX(modified) AS modified
		FROM `tabTIQN Payslip`
		GROUP BY period, status
		ORDER BY period DESC
		""",
		as_dict=True,
	)
	periods = {}
	for r in rows:
		p = periods.setdefault(
			r.period,
			{"period": r.period, "Draft": 0, "Published": 0, "Withdrawn": 0, "net_total": 0, "published_on": None, "viewed": 0,
			 "withdrawn_viewed": 0, "final_settlement": 0},
		)
		p[r.status] = r.n
		p["final_settlement"] += cint(r.final_settlement)
		if r.status == "Published":
			# Employees who opened their (currently published) payslip at least once.
			p["viewed"] = cint(r.viewed)
		elif r.status == "Withdrawn":
			# Shown before "Delete Payslips": withdrawn payslips some employees already read.
			p["withdrawn_viewed"] = cint(r.viewed)
		p["net_total"] += flt(r.net_total)
		if r.published_on and (not p["published_on"] or r.published_on > p["published_on"]):
			p["published_on"] = r.published_on
	out = []
	for p in periods.values():
		p["from_date"], p["to_date"] = period_bounds(p["period"])
		out.append(p)
	return out


@frappe.whitelist(methods=["GET"])
def get_config_errors():
	"""Shown on the Payslip Upload page: the period HR is about to compute (the last one
	that ended) has no Payslip Structure - the template and the upload would refuse it."""
	_check_role()
	period = shift_period(period_of_date(today()), -1)
	if pss.structure_for_period(period):
		return []
	return [_("No Payslip Structure covers period {0} yet. Create one before downloading the template.").format(period)]


@frappe.whitelist(methods=["POST"])
def publish_period(period):
	"""Publish every Draft / Withdrawn payslip of `period`, then ONE company-wide
	notification for the period (shown only to employees with a published payslip).
	Final settlements (left during the period, paid separately) alone send none: HR
	tells those people directly; the notification belongs to the monthly payroll."""
	_check_role()
	rows = frappe.get_all(
		"TIQN Payslip",
		filters={"period": period, "status": ["in", ["Draft", "Withdrawn"]]},
		fields=["name", "final_settlement"],
	)
	names = [r.name for r in rows]
	expected = pss.structure_for_period(period)
	stale = frappe.get_all(
		"TIQN Payslip",
		filters={"name": ("in", names or [""]), "structure": ("!=", expected or "")},
		pluck="employee",
	)
	if stale:
		frappe.throw(
			_("{0} payslip(s) of {1} do not use its Payslip Structure {2} (e.g. {3}). Upload them again.").format(
				len(stale), period, expected or "-", ", ".join(stale[:5])
			)
		)
	monthly = sum(1 for r in rows if not r.final_settlement)
	now = now_datetime()
	label = period_label(period)
	for name in names:
		doc = frappe.get_doc("TIQN Payslip", name)
		doc.status = "Published"
		doc.flags.portal_status_change = True
		doc.published_on = now
		doc.published_by = frappe.session.user
		doc.save(ignore_permissions=True)
	if monthly:
		notify_all("PAYSLIP_PUBLISHED", "Payslip {0}", "has been published", title_args=(label,), target_id=period)
	return {"published": len(names), "final_settlement": len(names) - monthly, "notified": bool(monthly)}


@frappe.whitelist(methods=["POST"])
def withdraw_period(period):
	"""Hide every Published payslip of `period` from the Mini App (to fix and republish)."""
	_check_role()
	names = frappe.get_all("TIQN Payslip", filters={"period": period, "status": "Published"}, pluck="name")
	for name in names:
		doc = frappe.get_doc("TIQN Payslip", name)
		doc.status = "Withdrawn"
		doc.flags.portal_status_change = True
		doc.save(ignore_permissions=True)
	# A "payslip published" pointing at a withdrawn payslip only leads to "not found".
	# Publishing again creates a fresh one (shown as new).
	for n in frappe.get_all(
		"TIQN Portal Notification", filters={"type": "PAYSLIP_PUBLISHED", "target_id": period}, pluck="name"
	):
		frappe.delete_doc("TIQN Portal Notification", n, ignore_permissions=True, force=True)
	return {"withdrawn": len(names)}


@frappe.whitelist(methods=["POST"])
def delete_period(period):
	"""Delete every Draft AND Withdrawn payslip of a period (user, 28/10/2026: a wrong upload,
	or a withdrawn period to redo from scratch). Published ones are never touched - withdraw
	them first. Frappe keeps each deleted payslip in Deleted Document (System Manager can
	restore it). Their notifications went with Withdraw. Over QUEUE_THRESHOLD in the background."""
	_check_role()
	period = validate_period(period)
	rows = frappe.get_all(
		"TIQN Payslip",
		filters={"period": period, "status": ["in", ["Draft", "Withdrawn"]]},
		fields=["name", "status", "first_viewed_on"],
	)
	names = [r.name for r in rows]
	result = {
		"deleted": len(names),
		"withdrawn": sum(1 for r in rows if r.status == "Withdrawn"),
		"viewed": sum(1 for r in rows if r.first_viewed_on),
	}
	if len(names) > QUEUE_THRESHOLD and not frappe.flags.in_test:
		frappe.enqueue(_delete_payslips, queue="long", timeout=1800, names=names, period=period, user=frappe.session.user)
		return {**result, "queued": True}
	_delete_payslips(names, period, frappe.session.user)
	return {**result, "queued": False}


def _delete_payslips(names, period, user):
	for name in names:
		frappe.delete_doc("TIQN Payslip", name, ignore_permissions=True)
	if frappe.flags.in_test:
		return
	frappe.db.commit()
	frappe.publish_realtime("payslip_delete_done", {"period": period, "deleted": len(names)}, user=user, after_commit=True)
