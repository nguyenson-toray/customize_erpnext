"""End-to-end check of the Employee Portal (Zalo Mini App) backend - M1.

Runs against the real site and rolls everything back. Covers the contract
(API_CONTRACT.md) AND the security rules: one employee's token must never reach
another employee's data, whatever the client sends.

Run:
    cd ~/frappe-bench/sites && ../env/bin/python -c "
    import frappe; frappe.init('erp.tiqn.local'); frappe.connect(); frappe.set_user('Administrator')
    exec(open('../apps/customize_erpnext/customize_erpnext/employee_portal/test_employee_portal.py').read())"

🔴 NOTHING here may commit, and endpoint errors must not roll back: the portal
wrapper calls frappe.db.rollback() on every ok:false (correct in production), which
would wipe this script's fixtures mid-run. Both are patched to no-ops for the run
and restored before the final rollback.
⚠ Assertions are scoped to records this script creates (zalo ids "unit-<uuid>",
periods 2030-xx) - the test team uses TIQN-0148 / TIQN-1405 on the same database.
"""

import io
import uuid
from unittest import mock

import frappe
from frappe.utils import add_days, add_to_date, flt, formatdate, getdate, now_datetime, today

import customize_erpnext.api.employee_portal as api
from customize_erpnext.api.employee_portal import auth as auth_mod
from customize_erpnext.api.employee_portal.core import hash_token
from customize_erpnext.employee_portal import payslip_admin, utils
from customize_erpnext.utils import zalo_graph

ok = fail = 0
RUN = uuid.uuid4().hex[:8]
EMP_A, EMP_B = "TIQN-0148", "TIQN-1405"

_REAL_COMMIT, _REAL_ROLLBACK = frappe.db.commit, frappe.db.rollback
frappe.db.commit = lambda *a, **k: None
frappe.flags.in_test = True


def check(label, got, want):
	global ok, fail
	if got == want:
		ok += 1
		print(f"  ok   {label}")
	else:
		fail += 1
		print(f"  FAIL {label}: got {got!r}, want {want!r}")


def call(fn, token=None, version=None, **kwargs):
	"""Call an endpoint the way an HTTP request would: headers in, envelope out."""
	headers = {}
	if token:
		headers["X-Portal-Token"] = token
	if version:
		headers["X-Portal-Client-Version"] = version

	def fake_header(name, default=None):
		return headers.get(name, default)

	lang = frappe.local.lang
	with mock.patch.object(frappe, "get_request_header", side_effect=fake_header), mock.patch.object(
		frappe.db, "rollback", lambda *a, **k: None
	):
		result = fn(**kwargs)
	frappe.set_user("Administrator")
	# The portal answers in Vietnamese (core sets frappe.local.lang = "vi"); a real request
	# ends there - the desk checks after this one read English messages again.
	frappe.local.lang = lang
	return result


def refused(fn):
	"""Saved/refused - kept as refused_fn too: a later section reuses the name `refused`."""
	try:
		fn()
		return "saved"
	except frappe.ValidationError:
		frappe.clear_messages()
		return "refused"


refused_fn = refused


def code(res):
	return None if res.get("ok") else res["error"]["code"]


def phone_of(emp):
	"""The number the employee signs in with: HR phone, else the Zalo phone column (HR may
	move an employee's number there while testing - TIQN-0148 on 24/10)."""
	cell, zalo = frappe.db.get_value("Employee", emp, ["cell_number", "custom_cell_phone_zalo"])
	return cell or zalo


# HR may also clear a test employee's phones (TIQN-1405 on 26/10): give it a throwaway
# Zalo phone for this run - rolled back at the end like everything else.
for _i, _emp in enumerate((EMP_A, EMP_B), start=1):
	if not phone_of(_emp):
		_phone = f"0999{_i}{int(RUN, 16) % 100000:05d}"
		assert not utils.find_employees_by_phone(_phone), _phone
		frappe.db.set_value("Employee", _emp, "custom_cell_phone_zalo", _phone, update_modified=False)


# ---------------------------------------------------------------------------
print("\n=== utils ===")
for raw in ("0962200054", "84962200054", "+84 962 200 054", "962200054"):
	check(f"phone_key({raw})", utils.phone_key(raw), "962200054")
check("phone_key too short", utils.phone_key("12345"), None)
check("mask_phone", utils.mask_phone("84962200054"), "09•• ••• 054")
check("period_bounds 2026-08", tuple(str(d) for d in utils.period_bounds("2026-08")), ("2026-07-26", "2026-08-25"))
check("period_bounds 2026-01", tuple(str(d) for d in utils.period_bounds("2026-01")), ("2025-12-26", "2026-01-25"))
check("period_of_date 25/08", utils.period_of_date("2026-08-25"), "2026-08")
check("period_of_date 26/08", utils.period_of_date("2026-08-26"), "2026-09")
check("period_of_date 26/12", utils.period_of_date("2026-12-26"), "2027-01")
check("shift_period", utils.shift_period("2026-01", -1), "2025-12")
for raw, want in (("10-2026", "2026-10"), ("10/2026", "2026-10"), ("2026-10", "2026-10"), ("13-2026", None), ("", None)):
	check(f"parse_period_cell({raw!r})", utils.parse_period_cell(raw), want)
import datetime as _dt

check("parse_period_cell(date)", utils.parse_period_cell(_dt.datetime(2026, 10, 1)), "2026-10")
check("initials", utils.initials("Nguyễn Văn An"), "NA")
check("employee_access Active", utils.employee_access({"status": "Active"})[0], "ACTIVE")
check(
	"employee_access Left in window ('Left ' with space)",
	utils.employee_access({"status": "Left ", "relieving_date": add_days(today(), -10)})[0],
	"READ_ONLY",
)
check("employee_access Left expired", utils.employee_access({"status": "Left", "relieving_date": add_days(today(), -61)})[0], "EXPIRED")
check("read_only_days is 30", utils.setting_int("read_only_days", 0), 30)
check("employee_access Inactive", utils.employee_access({"status": "Inactive"})[0], "EXPIRED")


# ---------------------------------------------------------------------------
print("\n=== login decision table ===")
zid_a, zid_b = f"unit-{RUN}-a", f"unit-{RUN}-b"
r = auth_mod.complete_login(zid_a, "84" + utils.phone_key(phone_of(EMP_A)))
check("A: ACTIVE", r["result"], "ACTIVE")
check("A: employee", r["employee"]["employee_id"], EMP_A)
check("A: token issued", bool(r["token"]) and len(r["token"]) >= 40, True)
check("A: token stored only as hash", frappe.db.exists("TIQN Portal Session", {"token_hash": hash_token(r["token"])}) is not None, True)
check("A: token not stored in clear", frappe.db.exists("TIQN Portal Session", {"token_hash": r["token"]}), None)
check("A: masked phone", r["masked_phone"].count("•"), 5)
check("A: every contract key present", sorted(r), sorted(["result", "token", "expires_on", "zalo_user_id", "masked_phone", "employee", "access", "oa", "expired_on", "message", "link_request"]))
check("A: avatar_url absolute + encoded", (r["employee"]["avatar_url"] or "").startswith("http") and " " not in r["employee"]["avatar_url"], True)
check("avatar_url: none -> None", auth_mod.avatar_url(None), None)
check("avatar_url: private -> None", auth_mod.avatar_url("/private/files/x.jpg"), None)
check("avatar_url: missing file -> None", auth_mod.avatar_url("/files/employee_photos/__missing__.jpg"), None)
TOKEN_A = r["token"]
TOKEN_B = auth_mod.complete_login(zid_b, phone_of(EMP_B))["token"]

# Zalo display name (v2): stored from Zalo, kept when Zalo sends none.
auth_mod.complete_login(zid_a, None, "STORED_PHONE", zalo_name="Unit Zalo Name")
check("zalo_name stored", frappe.db.get_value("TIQN Portal Account", zid_a, "zalo_name"), "Unit Zalo Name")
auth_mod.complete_login(zid_a, None, "STORED_PHONE")
check("zalo_name kept when Zalo sends none", frappe.db.get_value("TIQN Portal Account", zid_a, "zalo_name"), "Unit Zalo Name")

# Phone priority (23/10): cell_number first, the second column only when no cell_number
# matches. custom_cell_phone_zalo may not exist yet -> tested through a stand-in column.
taken = {utils.phone_key(x) for x in frappe.get_all("Employee", filters={"cell_number": ["is", "set"]}, pluck="cell_number")}
stand_in = next(
	(
		e
		for e in frappe.get_all(
			"Employee", filters={"emergency_phone_number": ["is", "set"]}, fields=["name", "emergency_phone_number"], order_by="name"
		)
		if utils.phone_key(e.emergency_phone_number) and utils.phone_key(e.emergency_phone_number) not in taken
	),
	None,
)
saved_fields = utils.PHONE_FIELDS
try:
	utils.PHONE_FIELDS = ("cell_number", "custom_field_that_does_not_exist")
	check("phone: missing 2nd column is skipped", utils.find_employees_by_phone("0000000001"), [])
	if stand_in:
		utils.PHONE_FIELDS = ("cell_number", "emergency_phone_number")
		m = utils.find_employees_by_phone(stand_in.emergency_phone_number)
		check("phone: 2nd column used when no cell_number matches", (m[0].name, m[0].matched_field), (stand_in.name, "emergency_phone_number"))
	# Any employee with a Mobile Phone - the test employees may have none (HR moves numbers while testing).
	with_cell = next((e for e in (EMP_A, EMP_B) if frappe.db.get_value("Employee", e, "cell_number")), None) or frappe.db.get_value(
		"Employee", {"status": "Active", "cell_number": ["is", "set"]}, "name"
	)
	m = utils.find_employees_by_phone(frappe.db.get_value("Employee", with_cell, "cell_number"))
	check("phone: cell_number wins", (with_cell in [x.name for x in m], m[0].matched_field), (True, "cell_number"))
finally:
	utils.PHONE_FIELDS = saved_fields

# Settings: a dependent feature cannot stay on when its parent is off.
st = frappe.get_doc("Employee Portal Settings")
st.feature_attendance, st.feature_overtime, st.feature_explanation_create = 0, 1, 1
st.feature_leave_view, st.feature_leave_create = 0, 1
st.clear_dependent_features()
check("settings: attendance off clears overtime + explanation", (st.feature_overtime, st.feature_explanation_create), (0, 0))
check("settings: leave view off clears leave create", st.feature_leave_create, 0)

# Chat with C&B (23/10): 1-1 chat with the payroll staff member's ID by OA, else OA, else hidden.
st = frappe.get_doc("Employee Portal Settings")
for bad in ("34", "abc1234567890", "123456789"):
	st.cb_chat_id_by_oa = bad
	check(f"settings: ID by OA {bad!r} refused", refused(st.validate), "refused")
saved_chat = frappe.db.get_value("Employee Portal Settings", None, ["cb_chat_id_by_oa", "cb_chat_name", "oa_chat_enabled", "oa_id"], as_dict=True)
frappe.db.set_single_value("Employee Portal Settings", {"cb_chat_id_by_oa": "4295057266797901281", "cb_chat_name": "Ms. Unit - Payroll"})
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
check("me.chat = C&B person", call(api.me, token=TOKEN_A)["data"]["chat"], {"type": "user", "id": "4295057266797901281", "name": "Ms. Unit - Payroll"})
frappe.db.set_single_value("Employee Portal Settings", {"cb_chat_id_by_oa": None, "oa_chat_enabled": 1, "oa_id": "1234567890123456789"})
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
check("me.chat falls back to OA", call(api.me, token=TOKEN_A)["data"]["chat"]["type"], "oa")
frappe.db.set_single_value("Employee Portal Settings", {"oa_chat_enabled": 0})
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
check("me.chat None when neither", call(api.me, token=TOKEN_A)["data"]["chat"], None)
frappe.db.set_single_value("Employee Portal Settings", dict(saved_chat))
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
check("get_config does not expose the C&B ID", "cb_chat" in frappe.as_json(call(api.get_config)["data"]) or "4295057266797901281" in frappe.as_json(call(api.get_config)["data"]), False)
check("set_id_by_oa: placeholder refused", code(call(api.set_id_by_oa, token=TOKEN_A, id_by_oa="34")), "VALIDATION")
check("set_id_by_oa: stored on own account", (call(api.set_id_by_oa, token=TOKEN_A, id_by_oa="7664780689128284452")["ok"],
	frappe.db.get_value("TIQN Portal Account", zid_a, "id_by_oa")), (True, "7664780689128284452"))

r = auth_mod.complete_login(f"unit-{RUN}-nf", "0900000" + RUN[:3].translate(str.maketrans("abcdef", "123456")))
check("unknown phone: NOT_FOUND", r["result"], "NOT_FOUND")
check("NOT_FOUND: no token", r["token"], None)
check("NOT_FOUND: account kept for HR report", frappe.db.get_value("TIQN Portal Account", f"unit-{RUN}-nf", "last_result"), "NOT_FOUND")

# New account without a matching employee starts auto-disabled, unlocks on a matching phone.
acc_nf = frappe.get_doc("TIQN Portal Account", f"unit-{RUN}-nf")
check("NOT_FOUND new account: auto-disabled", (acc_nf.disabled, acc_nf.auto_disabled), (1, 1))
check("ACTIVE new account: not disabled", frappe.db.get_value("TIQN Portal Account", zid_a, ["disabled", "auto_disabled"]), (0, 0))
r = auth_mod.complete_login(f"unit-{RUN}-nf", None, "STORED_PHONE")
check("auto-disabled + no match: still NOT_FOUND (not DISABLED)", r["result"], "NOT_FOUND")
r = auth_mod.complete_login(f"unit-{RUN}-nf", phone_of(EMP_A))
check("auto-disabled + Zalo phone matches: ACTIVE", r["result"], "ACTIVE")
check("auto-disabled cleared after match", frappe.db.get_value("TIQN Portal Account", f"unit-{RUN}-nf", ["disabled", "auto_disabled"]), (0, 0))

zid_ad = f"unit-{RUN}-ad"
r = auth_mod.complete_login(zid_ad, None, "STORED_PHONE")
check("no phone: NO_PHONE (not DISABLED)", r["result"], "NO_PHONE")
check("NO_PHONE new account: auto-disabled", frappe.db.get_value("TIQN Portal Account", zid_ad, ["disabled", "auto_disabled"]), (1, 1))
acc_ad = frappe.get_doc("TIQN Portal Account", zid_ad)
acc_ad.phone = phone_of(EMP_B)
acc_ad.save(ignore_permissions=True)
check("HR assigns matching phone: enabled", (acc_ad.disabled, acc_ad.auto_disabled, acc_ad.employee), (0, 0, EMP_B))
check("HR-assigned phone then login: ACTIVE", auth_mod.complete_login(zid_ad, None, "STORED_PHONE")["result"], "ACTIVE")

acc_hr = frappe.new_doc("TIQN Portal Account")
acc_hr.zalo_user_id = f"unit-{RUN}-hr"
acc_hr.insert(ignore_permissions=True)
check("HR creates without phone: auto-disabled", (acc_hr.disabled, acc_hr.auto_disabled), (1, 1))
acc_hr.auto_disabled = 0  # as if HR ticked Disabled by hand
acc_hr.save(ignore_permissions=True)
acc_hr.phone = phone_of(EMP_B)
acc_hr.save(ignore_permissions=True)
check("hand-disabled stays disabled after matching phone", acc_hr.disabled, 1)
check("hand-disabled + match: DISABLED", auth_mod.complete_login(f"unit-{RUN}-hr", None, "STORED_PHONE")["result"], "DISABLED")
acc_hr.reload()
acc_hr.disabled = 0
acc_hr.save(ignore_permissions=True)
acc_hr.disabled = 1
acc_hr.save(ignore_permissions=True)
check("HR ticks Disabled by hand: not auto", acc_hr.auto_disabled, 0)

zid_d = f"unit-{RUN}-d"
auth_mod.complete_login(zid_d, phone_of(EMP_A))
frappe.db.set_value("TIQN Portal Account", zid_d, "disabled", 1)
r = auth_mod.complete_login(zid_d, phone_of(EMP_A))
check("disabled account: DISABLED", r["result"], "DISABLED")
check("DISABLED: no token", r["token"], None)

left_recent = frappe.db.sql(
	"""SELECT name, cell_number FROM tabEmployee WHERE TRIM(status)='Left' AND relieving_date > %s
	   AND IFNULL(cell_number,'')!='' ORDER BY relieving_date DESC LIMIT 20""",
	add_days(today(), -25),
	as_dict=True,
)


def owns_phone(e):
	matches = utils.find_employees_by_phone(e.cell_number)
	return bool(matches) and matches[0].name == e.name


left_recent = [e for e in left_recent if owns_phone(e)]
if left_recent:
	r = auth_mod.complete_login(f"unit-{RUN}-ro", left_recent[0].cell_number)
	check("recently left: READ_ONLY", r["result"], "READ_ONLY")
	check("READ_ONLY: access.mode", r["access"]["mode"], "READ_ONLY")
	check("recently left: new account NOT auto-disabled", frappe.db.get_value("TIQN Portal Account", f"unit-{RUN}-ro", "disabled"), 0)
else:
	print("  skip READ_ONLY (no recently-left employee with a unique phone)")

left_old = frappe.db.sql(
	"""SELECT name, cell_number FROM tabEmployee WHERE TRIM(status)='Left' AND relieving_date < %s
	   AND IFNULL(cell_number,'')!='' ORDER BY relieving_date DESC LIMIT 50""",
	add_days(today(), -90),
	as_dict=True,
)
left_old = [e for e in left_old if owns_phone(e)]
if left_old:
	r = auth_mod.complete_login(f"unit-{RUN}-ex", left_old[0].cell_number)
	check("left long ago: EXPIRED", r["result"], "EXPIRED")
	check("EXPIRED: expired_on set", bool(r["expired_on"]), True)
	check("left long ago: new account auto-disabled", frappe.db.get_value("TIQN Portal Account", f"unit-{RUN}-ex", ["disabled", "auto_disabled"]), (1, 1))
	r = auth_mod.complete_login(f"unit-{RUN}-ex", None, "STORED_PHONE")
	check("auto-disabled + left long ago: still EXPIRED (not DISABLED)", r["result"], "EXPIRED")


print("\n=== login endpoint with Zalo mocked ===")
with mock.patch.object(zalo_graph, "get_phone", return_value="84" + utils.phone_key(phone_of(EMP_A))), mock.patch.object(
	zalo_graph, "get_user_id", return_value=(f"unit-{RUN}-http", "Test")
):
	r = call(api.login, access_token="at", phone_token="pt")
check("login: ok envelope", r["ok"], True)
check("login: ACTIVE", r["data"]["result"], "ACTIVE")
check("login: server_time", bool(r.get("server_time")), True)
r = call(api.login, access_token="", phone_token="pt")
check("login without access_token", (code(r), r["error"]["field"]), ("ZALO_TOKEN_INVALID", "access_token"))

with mock.patch.object(zalo_graph, "_secret", return_value="s"), mock.patch.object(
	zalo_graph, "_get", return_value={"error": 119, "message": "code has already been used"}
):
	check("Zalo 119 -> ZALO_TOKEN_USED", code(call(api.login, access_token="at", phone_token="pt")), "ZALO_TOKEN_USED")
with mock.patch.object(zalo_graph, "_secret", return_value="s"), mock.patch.object(
	zalo_graph, "_get", return_value={"error": 118, "message": "code is invalid"}
):
	check("Zalo 118 -> ZALO_SERVER_CONFIG", code(call(api.login, access_token="at", phone_token="pt")), "ZALO_SERVER_CONFIG")
with mock.patch.dict(frappe.conf, {"zalo_app_employee_portal_secret_key": ""}):
	check("no secret configured -> ZALO_SERVER_CONFIG", code(call(api.login, access_token="at", phone_token="pt")), "ZALO_SERVER_CONFIG")


# ---------------------------------------------------------------------------
print("\n=== login without a phone token / HR-assigned phone ===")
zid_np = f"unit-{RUN}-np"
with mock.patch.object(zalo_graph, "get_user_id", return_value=(zid_np, "Test")), mock.patch.object(zalo_graph, "get_phone") as gp:
	r = call(api.login, access_token="at")
	check("no phone_token: Zalo phone API not called", gp.called, False)
check("unknown user, no phone -> NO_PHONE", r["data"]["result"], "NO_PHONE")
check("NO_PHONE returns own zalo id for HR", r["data"]["zalo_user_id"], zid_np)
check("NO_PHONE: no token", r["data"]["token"], None)
check("NO_PHONE account recorded", frappe.db.get_value("TIQN Portal Account", zid_np, "last_result"), "NO_PHONE")

acc = frappe.get_doc("TIQN Portal Account", zid_np)
acc.phone = "84" + utils.phone_key(phone_of(EMP_A))
acc.save(ignore_permissions=True)
check("HR assigns phone: normalised", acc.phone, "0" + utils.phone_key(phone_of(EMP_A)))
check("HR assigns phone: source Manual", acc.phone_source, "Manual")
check("HR assigns phone: employee derived", acc.employee, EMP_A)
try:
	acc.phone = "0900000" + RUN[:3].translate(str.maketrans("abcdef", "123456"))
	acc.save(ignore_permissions=True)
	check("HR phone that matches nobody refused", "saved", "refused")
except frappe.ValidationError:
	check("HR phone that matches nobody refused", "refused", "refused")
	acc.reload()

with mock.patch.object(zalo_graph, "get_user_id", return_value=(zid_np, "Test")):
	r = call(api.login, access_token="at")
check("stored phone -> ACTIVE", (r["data"]["result"], r["data"]["employee"]["employee_id"]), ("ACTIVE", EMP_A))
tok_np = r["data"]["token"]
check("session login_method STORED_PHONE", frappe.db.get_value("TIQN Portal Session", {"token_hash": hash_token(tok_np)}, "login_method"), "STORED_PHONE")
check("me shows login_method", call(api.me, token=tok_np)["data"]["login_method"], "STORED_PHONE")

acc = frappe.get_doc("TIQN Portal Account", zid_np)
acc.phone = phone_of(EMP_B)
acc.save(ignore_permissions=True)
r = call(api.me, token=tok_np)
check("HR re-assigns phone -> old session ends", (code(r), r["error"]["details"].get("reason")), ("SESSION_INVALID", "PHONE_CHANGED"))

with mock.patch.object(zalo_graph, "get_user_id", return_value=(zid_np, "Test")), mock.patch.object(
	zalo_graph, "get_phone", return_value="84" + utils.phone_key(phone_of(EMP_A))
):
	r = call(api.login, access_token="at", phone_token="pt")
check("Zalo phone overrides HR phone", (r["data"]["employee"]["employee_id"], frappe.db.get_value("TIQN Portal Account", zid_np, "phone_source")), (EMP_A, "Zalo"))
check("session login_method ZALO_PHONE", frappe.db.get_value("TIQN Portal Session", {"token_hash": hash_token(r["data"]["token"])}, "login_method"), "ZALO_PHONE")


print("\n=== dev_login (simulator) ===")
zid_dev = f"unit-{RUN}-dev"
with mock.patch.dict(frappe.conf, {"employee_portal_dev_login": 0}):
	check("dev_login off -> FEATURE_DISABLED", code(call(api.dev_login, zalo_user_id=zid_dev)), "FEATURE_DISABLED")
	check("get_config.dev_login false", call(api.get_config)["data"]["dev_login"], False)
with mock.patch.dict(frappe.conf, {"employee_portal_dev_login": 1}):
	check("get_config.dev_login true", call(api.get_config)["data"]["dev_login"], True)
	r = call(api.dev_login, zalo_user_id=zid_dev)
	check("unknown id -> DEV_NOT_ALLOWED", (r["data"]["result"], r["data"]["token"]), ("DEV_NOT_ALLOWED", None))
	check("unknown id recorded for HR", frappe.db.exists("TIQN Portal Account", zid_dev) is not None, True)
	acc = frappe.get_doc("TIQN Portal Account", zid_dev)
	acc.phone = phone_of(EMP_A)
	acc.save(ignore_permissions=True)
	check("phone but not ticked -> still DEV_NOT_ALLOWED", call(api.dev_login, zalo_user_id=zid_dev)["data"]["result"], "DEV_NOT_ALLOWED")
	frappe.db.set_value("TIQN Portal Account", zid_dev, "allow_dev_login", 1)
	r = call(api.dev_login, zalo_user_id=zid_dev)
	check("ticked + phone -> ACTIVE", r["data"]["result"], "ACTIVE")
	tok_dev = r["data"]["token"]
	sess = frappe.db.get_value("TIQN Portal Session", {"token_hash": hash_token(tok_dev)}, ["login_method", "expires_on"], as_dict=True)
	check("dev session marked DEV", sess.login_method, "DEV")
	check("dev session lasts 1 day", sess.expires_on <= add_to_date(now_datetime(), days=1, minutes=1), True)
	check("dev session works", call(api.me, token=tok_dev)["ok"], True)
	check("dev_login cannot pick another account", call(api.dev_login, zalo_user_id=zid_np)["data"]["result"], "DEV_NOT_ALLOWED")
with mock.patch.dict(frappe.conf, {"employee_portal_dev_login": 0}):
	r = call(api.me, token=tok_dev)
	check("dev flag turned off -> dev session ends", code(r), "SESSION_INVALID")


# ---------------------------------------------------------------------------
print("\n=== sessions ===")
r = call(api.me, token=TOKEN_A)
check("me: ok", r["ok"], True)
check("me: employee", r["data"]["employee"]["employee_id"], EMP_A)
check("me: department suffix stripped", " - TIQN" in (r["data"]["employee"]["department"] or ""), False)
check("me without token", code(call(api.me)), "SESSION_REQUIRED")
check("me with garbage token", code(call(api.me, token="nope")), "SESSION_INVALID")
check("old client", code(call(api.me, token=TOKEN_A, version="0.9.0")), "CLIENT_TOO_OLD")
check("current client", call(api.me, token=TOKEN_A, version="1.0.0")["ok"], True)

tok_exp = auth_mod.complete_login(f"unit-{RUN}-e", phone_of(EMP_A))["token"]
frappe.db.set_value("TIQN Portal Session", {"token_hash": hash_token(tok_exp)}, "expires_on", add_to_date(now_datetime(), hours=-1))
r = call(api.me, token=tok_exp)
check("expired session", (code(r), r["error"]["details"].get("reason")), ("SESSION_INVALID", "EXPIRED"))

tok_out = auth_mod.complete_login(f"unit-{RUN}-o", phone_of(EMP_A))["token"]
check("logout ok", call(api.logout, token=tok_out)["ok"], True)
r = call(api.me, token=tok_out)
check("after logout", (code(r), r["error"]["details"].get("reason")), ("SESSION_INVALID", "REVOKED"))

tok_dis = auth_mod.complete_login(f"unit-{RUN}-x", phone_of(EMP_A))["token"]
acc = frappe.get_doc("TIQN Portal Account", f"unit-{RUN}-x")
acc.disabled = 1
acc.save(ignore_permissions=True)
r = call(api.me, token=tok_dis)
check("disabled after login revokes session", code(r), "SESSION_INVALID")

tok_ph = auth_mod.complete_login(f"unit-{RUN}-p", phone_of(EMP_B))["token"]
real_phone_b = phone_of(EMP_B)
phone_field_b = "cell_number" if frappe.db.get_value("Employee", EMP_B, "cell_number") else "custom_cell_phone_zalo"
frappe.db.set_value("Employee", EMP_B, phone_field_b, "0911" + RUN[:6].translate(str.maketrans("abcdef", "123456")), update_modified=False)
r = call(api.me, token=tok_ph)
check("HR changed the phone", (code(r), (r.get("error") or {}).get("details", {}).get("reason")), ("SESSION_INVALID", "PHONE_CHANGED"))
frappe.db.set_value("Employee", EMP_B, phone_field_b, real_phone_b, update_modified=False)

r = call(api.set_oa_follow, token=TOKEN_A, action="DISMISSED")
check("set_oa_follow DISMISSED", r["data"], {"followed": False, "show_invite": False})
check("set_oa_follow bad action", code(call(api.set_oa_follow, token=TOKEN_A, action="x")), "VALIDATION")


# ---------------------------------------------------------------------------
print("\n=== config ===")
r = call(api.get_config)
check("get_config without token", r["ok"], True)
check("get_config features.payslip", r["data"]["features"]["payslip"], True)
check("get_config policy url absolute", r["data"]["legal"]["privacy_policy_url"].startswith("http"), True)


# ---------------------------------------------------------------------------
print("\n=== payslips ===")
from customize_erpnext.employee_portal import payslip_structure as pss
STRUCT, SLINES = pss.period_structure_lines("2030-03")  # structure of the test periods (2030-xx)
check("payslip: no child table any more", frappe.get_meta("TIQN Payslip").get_table_fields(), [])
# Phase 2 (28/10/2026): the per-line sc_* fields / columns and Settings > Payslip Lines are gone.
check("payslip: no sc_* field or column left", (
	[f.fieldname for f in frappe.get_meta("TIQN Payslip").fields if f.fieldname.startswith("sc_")],
	[c for c in frappe.db.get_table_columns("TIQN Payslip") if c.startswith("sc_")]), ([], []))
check("settings: Payslip Lines gone", (frappe.get_meta("Employee Portal Settings").has_field("payslip_lines"),
	frappe.db.exists("DocType", "TIQN Portal Payslip Line")), (False, None))


def code_key(abbr):
	return [int(p) for p in str(abbr).split(".")]


def _fake_amounts(employee, period):
	"""(amounts, days / hours values) by code, deterministic per (employee, period);
	c=a+b, f=c+d+e, j=f-g+h+i and the group totals (1 = 1.2, 2 = 2.1+2.2+2.3) hold exactly."""
	seed = sum(ord(ch) for ch in employee + period)
	basic = 7_000_000 + (seed % 7) * 500_000
	ot_normal = 6 + seed % 9 * 0.5
	ot_weekend = (seed % 3) * 4
	ot_holiday = 0
	hourly = round(basic / 26 / 8)
	q = {"1.2": 26 - seed % 2, "2.1": ot_normal, "2.2": ot_weekend, "2.3": ot_holiday, "7.4": 0, "7.5": 0}
	a = {
		"1.1": 26,
		"1.2": basic,
		"2.1": round(hourly * ot_normal * 1.5),
		"2.2": round(hourly * ot_weekend * 2),
		"2.3": round(hourly * ot_holiday * 3),
		"3.2": 500_000,
		"3.3": 0,
		"3.4": 0,
		"3.5": 300_000,
		"3.6": 200_000,
		"3.7": 0,
		"3.8": 150_000 if seed % 2 else 0,
		"3.9": 0,
		"3.10": 0,
		"4.1": 500_000 if seed % 4 else 0,
		"4.2": 0,
		"4.3": 300_000,
		"4.4": 0,
		"4.5": 0,
		"4.6": 0,
		"4.7": 0,
		"6.1": round(basic * 0.08),
		"6.2": round(basic * 0.015),
		"6.3": round(basic * 0.01),
		"6.4": min(round(basic * 0.01), 234_000),
		"6.5": 0,
		"6.6": 0,
		"7.1": 650_000,
		"7.2": 0,
		"7.3": 0,
		"7.4": 0,
		"7.5": 0,
		"7.6": 0,
		"7.7": 0,
		"8": 150_000 if seed % 5 == 0 else 0,
	}
	a["1"] = a["1.2"]
	a["2"] = a["2.1"] + a["2.2"] + a["2.3"]
	a["3"] = a["1"] + a["2"]
	a["3.1"] = sum(a[f"3.{i}"] for i in range(2, 11))
	a["4"] = sum(a[f"4.{i}"] for i in range(1, 8))
	a["5"] = a["3"] + a["3.1"] + a["4"]
	a["6"] = sum(a[f"6.{i}"] for i in range(1, 7))
	a["7"] = sum(a[f"7.{i}"] for i in range(1, 8))
	a["9"] = a["5"] - a["6"] + a["7"] + a["8"]
	return a, q


def _snap(a, q, lines=None):
	"""lines_json of amounts `a` + days / hours `q` (a days-only line, 1.1, takes its days from `a`)."""
	lines = lines or SLINES
	values = {l.code: {"amount": a.get(l.code, 0), "quantity": (a if l.unit != "VND" else q).get(l.code, 0)} for l in lines}
	return pss.dump_snapshot(pss.make_snapshot(lines, values))


def make_payslip(emp, period, status="Published"):
	a, q = _fake_amounts(emp, period)
	doc = frappe.get_doc(
		{
			"doctype": "TIQN Payslip",
			"employee": emp,
			"employee_name": "UNIT",
			"period": period,
			"status": status,
			"published_on": now_datetime() if status == "Published" else None,
			"import_batch": f"UNIT-{RUN}",
			"lines_json": _snap(a, q),
		}
	)
	doc.flags.portal_status_change = True
	doc.insert(ignore_permissions=True)
	return doc, a


doc_a, amounts_a = make_payslip(EMP_A, "2030-03")
qty_a = _fake_amounts(EMP_A, "2030-03")[1]
make_payslip(EMP_A, "2030-04", status="Draft")
make_payslip(EMP_B, "2030-05")

r = call(api.get_payslip_periods, token=TOKEN_A, year=2030)
periods = [i["period"] for i in r["data"]["items"]]
check("periods: only own Published", periods, ["2030-03"])
check("periods: no amounts anywhere", "net" in frappe.as_json(r["data"]) or "amount" in frappe.as_json(r["data"]), False)
check("periods: is_new before viewing", r["data"]["items"][0]["is_new"], True)

check("config: payslip.show_zero_lines = Settings (1.9.0)", call(api.get_config)["data"]["payslip"]["show_zero_lines"], bool(utils.get_settings().show_zero_lines))
r = call(api.get_payslip, token=TOKEN_A, period="2030-03", include_zero=0)
check("payslip: ok", r["ok"], True)
d = r["data"]
check("payslip: net_pay = component 9", d["net_pay"], round(amounts_a["9"]))
check("payslip: own employee", d["employee"]["employee_id"], EMP_A)
keys = [s["key"] for s in d["sections"]]
check("payslip: section order", keys, ["EARNINGS", "DEDUCTIONS", "OTHERS", "ADJUSTMENT", "ATTENDANCE"])
earn = d["sections"][0]
check("payslip: earnings total = component 5", earn["total"], round(amounts_a["5"]))
check("payslip: deductions sign", d["sections"][1]["sign"], -1)
check("payslip: deductions total = component 6", d["sections"][1]["total"], round(amounts_a["6"]))
allowance = next(l for l in earn["lines"] if l["code"] == "3.1")
check("payslip: 3.1 is a subtotal", allowance["is_total"], True)
check("payslip: 3.3 (0đ) hidden with include_zero=0", any(c["code"] == "3.3" for c in allowance["children"]), False)
check("payslip: 3.2 under 3.1", any(c["code"] == "3.2" for c in allowance["children"]), True)
att = d["sections"][-1]
att_12 = next(l for l in att["lines"] if l["code"] == "1.2")
check("payslip: sections (deprecated) still show 1.2 days in ATTENDANCE", (att_12["unit"], att_12["amount"]), ("DAY", qty_a["1.2"]))
# rows (1.3.0): the paper payslip line by line.
rows_by_no = {x["no"]: x for x in d["rows"]}
check("rows: hidden zero lines, headers kept", ("3.3" in rows_by_no, "3.1" in rows_by_no, "2.3" in rows_by_no), (False, True, False))
check("rows: 1 = header, formula a, level 0", (rows_by_no["1"]["is_header"], rows_by_no["1"]["formula"], rows_by_no["1"]["level"]), (True, "a", 0))
check("rows: 1.1 days only", {k: rows_by_no["1.1"][k] for k in ("amount", "quantity", "quantity_unit", "level")},
	{"amount": None, "quantity": 26, "quantity_unit": "DAY", "level": 1})
check("rows: 1.2 days + amount", (rows_by_no["1.2"]["quantity"], rows_by_no["1.2"]["quantity_unit"], rows_by_no["1.2"]["amount"]),
	(qty_a["1.2"], "DAY", amounts_a["1.2"]))
check("rows: 2.1 hours + amount", (rows_by_no["2.1"]["quantity"], rows_by_no["2.1"]["quantity_unit"], rows_by_no["2.1"]["amount"]),
	(qty_a["2.1"], "HOUR", amounts_a["2.1"]))
check("rows: 3.1 header level 0, 3.2 under it", (rows_by_no["3.1"]["is_header"], rows_by_no["3.2"]["level"], rows_by_no["3.2"]["is_header"]), (True, 1, False))
check("rows: 6 header, amount positive", (rows_by_no["6"]["is_header"], rows_by_no["6"]["amount"]), (True, round(amounts_a["6"])))
check("rows: 6.1 under 6", rows_by_no["6.1"]["level"], 1)
check("rows: plain amount line has no days/hours", (rows_by_no["6.1"]["quantity"], rows_by_no["6.1"]["quantity_unit"]), (None, None))
check("rows: formulas", (rows_by_no["3"]["formula"], rows_by_no["9"]["formula"], rows_by_no["6.2"]["formula"]), ("c=a+b", "j=f-g+h+i", None))
check("rows: display_label = No. + Vietnamese Label + (Formula) of the structure", (rows_by_no["1"]["display_label"], rows_by_no["1.2"]["display_label"], rows_by_no["9"]["display_label"]),
	("1 Lương cơ bản (a)", "1.2 Số ngày công thực tế làm việc", "9 Thực lĩnh (j=f-g+h+i)"))
check("rows: label = the text without number / letters (older app builds)", (rows_by_no["1"]["label"], rows_by_no["1.2"]["label"], rows_by_no["9"]["label"]),
	("Lương cơ bản", "Số ngày công thực tế làm việc", "Thực lĩnh"))
check("payslip: attendance total null", att["total"], None)
check("payslip: section-total line not repeated", any(l["code"] == "5" for l in earn["lines"]), False)
check("payslip: first_viewed_on set", bool(frappe.db.get_value("TIQN Payslip", doc_a.name, "first_viewed_on")), True)
check("payslip: View Log DocType gone", frappe.db.exists("DocType", "TIQN Payslip View Log"), None)
check("payslip: old view columns gone", {"viewed", "view_count", "last_viewed_on"} & set(frappe.db.get_table_columns("TIQN Payslip")), set())
v = frappe.db.get_value("TIQN Payslip", doc_a.name, ["first_viewed_on", "modified"], as_dict=True)
check("payslip: net_pay = the Net line", frappe.db.get_value("TIQN Payslip", doc_a.name, "net_pay"), round(amounts_a["9"]))
check("payslip: report columns from the lines", [frappe.db.get_value("TIQN Payslip", doc_a.name, f) for f in ("basic_pay", "gross_pay", "total_deductions", "actual_days", "ot_hours")],
	[amounts_a["1"], amounts_a["5"], amounts_a["6"], qty_a["1.2"], qty_a["2.1"] + qty_a["2.2"] + qty_a["2.3"]])
r = call(api.get_payslip, token=TOKEN_A, period="2030-03", include_zero=1)
allowance = next(l for l in r["data"]["sections"][0]["lines"] if l["code"] == "3.1")
check("payslip: include_zero shows 3.3", any(c["code"] == "3.3" for c in allowance["children"]), True)
check("rows: include_zero = every line in paper order", [x["no"] for x in r["data"]["rows"]], sorted([l.code for l in SLINES], key=code_key))
check("rows: zero line amount 0 (Mini App shows -)", next(x for x in r["data"]["rows"] if x["no"] == "3.3")["amount"], 0)
v2 = frappe.db.get_value("TIQN Payslip", doc_a.name, ["first_viewed_on", "modified"], as_dict=True)
check("payslip: 2nd view keeps first_viewed_on, modified untouched", (v2.first_viewed_on, v2.modified), (v.first_viewed_on, v.modified))
with mock.patch("customize_erpnext.api.employee_portal.payslip._mark_viewed") as mv:
	call(api.get_payslip, token=TOKEN_A, period="2030-03")
check("payslip: later views write nothing", mv.called, False)
r = call(api.get_payslip_periods, token=TOKEN_A, year=2030)
check("periods: is_new false after viewing", r["data"]["items"][0]["is_new"], False)

check("payslip of B with A's token -> NOT_FOUND", code(call(api.get_payslip, token=TOKEN_A, period="2030-05")), "NOT_FOUND")
check("Draft payslip -> NOT_FOUND", code(call(api.get_payslip, token=TOKEN_A, period="2030-04")), "NOT_FOUND")
check("bad period", code(call(api.get_payslip, token=TOKEN_A, period="2030-13")), "VALIDATION")
check("employee kwarg is ignored", call(api.get_payslip, token=TOKEN_A, period="2030-03", employee=EMP_B)["data"]["employee"]["employee_id"], EMP_A)

frappe.db.set_single_value("Employee Portal Settings", "feature_payslip", 0)
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
check("feature off -> FEATURE_DISABLED", code(call(api.get_payslip_periods, token=TOKEN_A)), "FEATURE_DISABLED")
frappe.db.set_single_value("Employee Portal Settings", "feature_payslip", 1)
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")

pub = frappe.get_doc("TIQN Payslip", doc_a.name)
_pl = frappe.parse_json(pub.lines_json)
_pl[0]["amount"] = 1
pub.lines_json = frappe.as_json(_pl)
try:
	pub.save(ignore_permissions=True)
	check("published payslip cannot be edited", "saved", "refused")
except frappe.ValidationError:
	check("published payslip cannot be edited", "refused", "refused")




# Data Import / hand edit cannot publish or withdraw (no notification, no published_on).
sneaky = frappe.get_doc(
	{"doctype": "TIQN Payslip", "employee": EMP_B, "period": "2030-09", "status": "Published", "import_batch": f"UNIT-{RUN}"}
)
check("new payslip with Status=Published refused", refused(lambda: sneaky.insert(ignore_permissions=True)), "refused")
draft_b = frappe.get_doc({"doctype": "TIQN Payslip", "employee": EMP_B, "period": "2030-09", "import_batch": f"UNIT-{RUN}"}).insert(
	ignore_permissions=True
)
draft_b.status = "Published"
check("Draft -> Published by editing refused", refused(lambda: draft_b.save(ignore_permissions=True)), "refused")
draft_b.reload()
check("hand-made payslip: the structure's lines at 0", (draft_b.structure, len(pss.snapshot_of(draft_b)), draft_b.net_pay), (STRUCT, len(SLINES), 0))


def _edit_line(doc, code_, **kw):
	lines_ = frappe.parse_json(doc.lines_json)
	next(l for l in lines_ if l["code"] == code_).update(kw)
	doc.lines_json = frappe.as_json(lines_)


_edit_line(draft_b, "9", amount=123)
check("Draft numbers editable", refused(lambda: draft_b.save(ignore_permissions=True)), "saved")
check("Draft edit: net_pay follows the Net line", frappe.db.get_value("TIQN Payslip", draft_b.name, "net_pay"), 123)
for code_, kw in (("1.2", {"quantity": -1}), ("1.2", {"quantity": 40}), ("2.1", {"quantity": 24 * 31 + 1}), ("1", {"amount": 10**11})):
	draft_b.reload()
	_edit_line(draft_b, code_, **kw)
	check(f"impossible number refused: {code_} {kw}", refused(lambda: draft_b.save(ignore_permissions=True)), "refused")
draft_b.reload()
_edit_line(draft_b, "9", amount=-500)
check("negative net salary allowed (HR decides)", refused(lambda: draft_b.save(ignore_permissions=True)), "saved")
check("Data Import off on TIQN Payslip (lines live in lines_json)", frappe.get_meta("TIQN Payslip").allow_import, 0)


# ---------------------------------------------------------------------------
print("\n=== upload / publish ===")
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
cols = []  # (header, code, is Unit/ĐVT) in template order
for l in SLINES:
	if l.quantity_unit:
		cols.append((payslip_admin.qty_header(l), l.code, True))
	cols.append((payslip_admin.amount_header(l), l.code, False))
comp_cols = [h for h, _code, _q in cols]


def vals(amounts_qty):
	a, q = amounts_qty
	return [(q if is_qty else a).get(code_, 0) for _h, code_, is_qty in cols]


ws.append(["Kỳ lương", "MSNV", "Họ tên", "dept", "custom_section", "custom_group"] + comp_cols)
good = _fake_amounts(EMP_A, "2030-07")
ws.append(["07-2030", EMP_A, frappe.db.get_value("Employee", EMP_A, "employee_name"), "Office", "HR-GA", "IT"] + vals(good))
bad_a, bad_q = _fake_amounts(EMP_B, "2030-07")
bad_a = dict(bad_a)
bad_a["9"] += 1000
bad_a["2.2"] += 1000  # 2 != 2.1 + 2.2 + 2.3
bad_formula = (bad_a, bad_q)
ws.append(["07-2030", EMP_B, "Tên Sai", "Office", "HR-GA", "IT"] + vals(bad_formula))
ws.append(["07-2030", "TIQN-9999", "Ghost", "", "", ""] + [0] * len(comp_cols))
ws.append(["07-2030", EMP_A, "", "", "", ""] + ["abc"] + [0] * (len(comp_cols) - 1))
buf = io.BytesIO()
wb.save(buf)
f = frappe.get_doc({"doctype": "File", "file_name": f"unit-{RUN}.xlsx", "is_private": 1, "content": buf.getvalue()}).insert(
	ignore_permissions=True
)
data = payslip_admin.parse_upload(f.file_url)
rows = {(r["row"]): r for r in data["rows"]}
check("upload: 4 rows read", len(data["rows"]), 4)
check("upload: good row clean", (rows[2]["errors"], rows[2]["warnings"]), ([], []))
check("upload: formulas never checked (HR's numbers)", [w for w in rows[3]["warnings"] if w[:4] in ("9 = ", "2 = ")], [])
check("upload: Unit/ĐVT values read", (rows[2]["values"]["1.2"]["quantity"], rows[2]["values"]["1.2"]["amount"]), (good[1]["1.2"], good[0]["1.2"]))
check("upload: structure of the file", data["structure"], STRUCT)
check("upload: template headers = all columns", data["missing_components"], [])
check("upload: name mismatch warning", any("Tên Sai" in w for w in rows[3]["warnings"]), True)
check("upload: unknown employee error", bool(rows[4]["errors"]), True)
check("upload: duplicate + non-number errors", len(rows[5]["errors"]) >= 2, True)
check("upload: periods", data["periods"], ["2030-07"])
throws_import = None
try:
	payslip_admin.import_upload(f.file_url, confirm_warnings=1)
except frappe.ValidationError:
	throws_import = "refused"
check("import refused while errors remain", throws_import, "refused")

wb2 = Workbook()
ws2 = wb2.active
ws2.append(["Kỳ lương", "MSNV", "Họ tên"] + comp_cols)
ws2.append(["07-2030", EMP_A, ""] + vals(good))
ws2.append(["07-2030", EMP_B, ""] + vals((dict(bad_a, **{"9": -5000}), bad_q)))  # negative net = a warning
buf2 = io.BytesIO()
wb2.save(buf2)
f2 = frappe.get_doc({"doctype": "File", "file_name": f"unit2-{RUN}.xlsx", "is_private": 1, "content": buf2.getvalue()}).insert(
	ignore_permissions=True
)
try:
	payslip_admin.import_upload(f2.file_url, confirm_warnings=0)
	check("import needs warning confirmation", "imported", "refused")
except frappe.ValidationError:
	check("import needs warning confirmation", "refused", "refused")
check("upload: negative net salary only warned (one warning, no error)",
	[(len(r["errors"]), len(r["warnings"])) for r in payslip_admin.parse_upload(f2.file_url)["rows"]], [(0, 0), (0, 1)])
res = payslip_admin.import_upload(f2.file_url, confirm_warnings=1)
check("import: 2 drafts", res["rows"], 2)
check("import: upload file deleted", frappe.db.exists("File", {"file_url": f2.file_url}), None)
check("import: status Draft", frappe.db.get_value("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"}, "status"), "Draft")
check("draft invisible in app", code(call(api.get_payslip, token=TOKEN_A, period="2030-07")), "NOT_FOUND")

res = payslip_admin.publish_period("2030-07")
check("publish: 2, notified", (res["published"], res["notified"]), (2, True))
check("published visible in app", call(api.get_payslip, token=TOKEN_A, period="2030-07")["ok"], True)
n = call(api.get_notifications, token=TOKEN_A)
mine = [i for i in n["data"]["items"] if i["target"]["id"] == "2030-07"]
check("notification created", (len(mine), mine[0]["type"] if mine else None), (1, "PAYSLIP_PUBLISHED"))
rows_0707 = frappe.get_all("TIQN Portal Notification", filters={"target_id": "2030-07"}, fields=["name", "audience"])
check("ONE company-wide row for the period (not one per employee)", [r.audience for r in rows_0707], ["All"])
check("company-wide row has no recipients", frappe.get_doc("TIQN Portal Notification", rows_0707[0].name).employees, [])
check("company-wide body in Vietnamese", mine[0]["body"], "đã được phát hành")
check("company-wide: unread for A before seeing", mine[0]["is_read"], False)
nb = call(api.get_notifications, token=TOKEN_B)
check("B sees the same row", [i["id"] for i in nb["data"]["items"] if i["target"]["id"] == "2030-07"], [mine[0]["id"]])

# Payslip broadcast only for employees who have a published payslip that period
# (2030-05: only B, published by make_payslip).
from customize_erpnext.employee_portal.notify import notify, notify_all

only_b = notify_all("PAYSLIP_PUBLISHED", "Payslip {0}", "has been published", title_args=("05/2030",), target_id="2030-05")
check("notify_all idempotent", notify_all("PAYSLIP_PUBLISHED", "Payslip {0}", target_id="2030-05").name, only_b.name)
ids_a = [i["id"] for i in call(api.get_notifications, token=TOKEN_A, page_size=100)["data"]["items"]]
ids_b = [i["id"] for i in call(api.get_notifications, token=TOKEN_B, page_size=100)["data"]["items"]]
check("payslip broadcast hidden without a payslip that period", (only_b.name in ids_a, only_b.name in ids_b), (False, True))

unread_b_before = call(api.get_notifications, token=TOKEN_B)["data"]["unread"]
r = call(api.mark_notifications_read, token=TOKEN_A, ids=mine[0]["id"])
again = [i for i in call(api.get_notifications, token=TOKEN_A)["data"]["items"] if i["id"] == mine[0]["id"]]
check("A marks the company-wide one read (own seen mark)", again[0]["is_read"], True)
check("A's seen mark stored on A's account only", (bool(frappe.db.get_value("TIQN Portal Account", zid_a, "notifications_seen_on")),
	frappe.db.get_value("TIQN Portal Account", zid_b, "notifications_seen_on")), (True, None))
check("B's unread not changed by A", call(api.get_notifications, token=TOKEN_B)["data"]["unread"], unread_b_before)

# Picked employees (Table MultiSelect): one row, several recipients; nobody else sees it.
def visible_ids(token):
	return [i["id"] for i in call(api.get_notifications, token=token, page_size=100)["data"]["items"]]


both = notify([EMP_A, EMP_B, EMP_A], "GENERAL", "Payslip {0}", title_args=("unit-both",))
check("notify to several: one row, duplicates dropped", ([r.employee for r in both.employees], both.audience), ([EMP_A, EMP_B], "Employees"))
check("several: both see it", (both.name in visible_ids(TOKEN_A), both.name in visible_ids(TOKEN_B)), (True, True))
b_note = notify(EMP_B, "GENERAL", "Payslip {0}", title_args=("unit",)).name
check("A does not see B's notification", b_note in visible_ids(TOKEN_A), False)
unread_b = call(api.get_notifications, token=TOKEN_B)["data"]["unread"]
call(api.mark_notifications_read, token=TOKEN_A, ids=b_note)
check("A cannot mark B's notification read", call(api.get_notifications, token=TOKEN_B)["data"]["unread"], unread_b)
call(api.mark_notifications_read, token=TOKEN_B, ids=b_note)
check("B marks own notification read", next(i for i in call(api.get_notifications, token=TOKEN_B, page_size=100)["data"]["items"] if i["id"] == b_note)["is_read"], True)
hr_doc = frappe.get_doc({"doctype": "TIQN Portal Notification", "audience": "Employees", "type": "GENERAL", "title": "x"})
check("Employees without anyone refused", refused(lambda: hr_doc.insert(ignore_permissions=True)), "refused")
hr_doc = frappe.get_doc({"doctype": "TIQN Portal Notification", "audience": "Employees", "type": "GENERAL", "title": "x", "employees": [{"employee": EMP_A}]}).insert(ignore_permissions=True)
check("Desk-made notification gets target screen from type", hr_doc.target_screen, "NONE")
call(api.mark_notifications_read, token=TOKEN_B, all=1)
check("mark all -> B has 0 unread", call(api.get_notifications, token=TOKEN_B)["data"]["unread"], 0)

check("re-import over Published refused", bool(payslip_admin.parse_upload(f.file_url)["rows"][0]["errors"]), True)
res = payslip_admin.withdraw_period("2030-07")
check("withdraw: 2", res["withdrawn"], 2)
check("withdrawn invisible in app", code(call(api.get_payslip, token=TOKEN_A, period="2030-07")), "NOT_FOUND")
check("withdraw removes the company-wide notification", frappe.db.exists("TIQN Portal Notification", mine[0]["id"]), None)

# Template layout (24/10, + Note 27/10): No. | Period | From | To | Pay Date | Employee | Name | Designation | Dept | Section | Group | Note.
wb3 = Workbook()
ws3 = wb3.active
ws3.append(list(payslip_admin.TEMPLATE_IDENT) + comp_cols)
ws3.append([1, "2030-07", getdate("2030-06-26"), "2030-07-25", None, EMP_A, "", "Any text", "Office", "HR-GA", "IT", None] + vals(good))
ws3.append([2, "2030-08", "2030-07-20", "2030-08-25", None, EMP_B, "", "", "", "", "", None] + vals(good))
ws3.append([3, "2030-08", "abc", None, None, EMP_A, "", "", "", "", "", None] + vals(good))
ws3.append([4, "2030-09", "2030-08-30", "2030-09-20", "2030-09-30", EMP_A, "", "", "", "", "", None] + vals(good))
buf3 = io.BytesIO()
wb3.save(buf3)
f3 = frappe.get_doc({"doctype": "File", "file_name": f"unit3-{RUN}.xlsx", "is_private": 1, "content": buf3.getvalue()}).insert(
	ignore_permissions=True
)
data3 = payslip_admin.parse_upload(f3.file_url)
check("template layout: no header errors", (data3["header_errors"], data3["missing_components"]), ([], []))
rows3 = {r["row"]: r for r in data3["rows"]}
check("template layout: matching Period From/To ok", (rows3[2]["errors"], rows3[2]["section"], rows3[2]["group"]), ([], "HR-GA", "IT"))
check("template layout: Period From before the period is an error", any("2030-08" in e for e in rows3[3]["errors"]), True)
check("template layout: unreadable Period From is an error", bool(rows3[4]["errors"]), True)
check("template layout: dates inside the period read", (rows3[5]["errors"], str(rows3[5]["period_from"]), str(rows3[5]["period_to"]), str(rows3[5]["pay_date"])),
	([], "2030-08-30", "2030-09-20", "2030-09-30"))
payslip_admin.discard_upload(f3.file_url)

# Re-importing a Withdrawn payslip keeps it Withdrawn (published once: view log + notification
# link to it); Publish takes it again.
wb4 = Workbook()
ws4 = wb4.active
ws4.append(list(payslip_admin.TEMPLATE_IDENT) + comp_cols)
ws4.append([1, "2030-07", None, None, None, EMP_A, "", "", "", "", "", None] + vals(good))
buf4 = io.BytesIO()
wb4.save(buf4)
f4 = frappe.get_doc({"doctype": "File", "file_name": f"unit4-{RUN}.xlsx", "is_private": 1, "content": buf4.getvalue()}).insert(
	ignore_permissions=True
)
old_batch = frappe.db.get_value("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"}, "import_batch")
row4 = payslip_admin.parse_upload(f4.file_url)["rows"][0]
check("re-import: overwrite warned (existing status + batch)", (row4["existing_status"], any(old_batch in w for w in row4["warnings"])), ("Withdrawn", True))
check("re-import: overwrite needs confirmation", refused_fn(lambda: payslip_admin.import_upload(f4.file_url, confirm_warnings=0)), "refused")
payslip_admin.import_upload(f4.file_url, confirm_warnings=1)
check("re-import withdrawn: stays Withdrawn", frappe.db.get_value("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"}, "status"), "Withdrawn")
wb5 = Workbook()
ws5 = wb5.active
ws5.append(list(payslip_admin.TEMPLATE_IDENT) + comp_cols)
ws5.append([1, "2030-07", None, getdate("2030-07-20"), None, EMP_A, "", "", "", "", "", None] + vals(good))
buf5 = io.BytesIO()
wb5.save(buf5)
f5 = frappe.get_doc({"doctype": "File", "file_name": f"unit5-{RUN}.xlsx", "is_private": 1, "content": buf5.getvalue()}).insert(ignore_permissions=True)
w5 = payslip_admin.parse_upload(f5.file_url)["rows"][0]["warnings"]
check("re-import: changed Period To shown (old and new)",
	any(formatdate("2030-07-25") in w and formatdate("2030-07-20") in w for w in w5), True)
payslip_admin.discard_upload(f5.file_url)
# Left mid-period but Period To in the file on/after the relieving date (monthly file downloaded
# before the relieving date was entered) -> warned; the controller would keep the file's value.
real_relieving = frappe.db.get_value("Employee", EMP_A, "relieving_date")
frappe.db.set_value("Employee", EMP_A, "relieving_date", "2030-07-20", update_modified=False)
left_warn = {}
for to_ in ("2030-07-25", "2030-07-20", "2030-07-19", None):
	wb6 = Workbook()
	ws6 = wb6.active
	ws6.append(list(payslip_admin.TEMPLATE_IDENT) + comp_cols)
	ws6.append([1, "2030-07", None, getdate(to_) if to_ else None, None, EMP_A, "", "", "", "", "", None] + vals(good))
	buf6 = io.BytesIO()
	wb6.save(buf6)
	f6 = frappe.get_doc({"doctype": "File", "file_name": f"unit6-{RUN}.xlsx", "is_private": 1, "content": buf6.getvalue()}).insert(ignore_permissions=True)
	left_warn[to_] = any("last working day" in w for w in payslip_admin.parse_upload(f6.file_url)["rows"][0]["warnings"])
	payslip_admin.discard_upload(f6.file_url)
frappe.db.set_value("Employee", EMP_A, "relieving_date", real_relieving, update_modified=False)
check("left mid-period: Period To >= relieving warned, < relieving / empty not",
	left_warn, {"2030-07-25": True, "2030-07-20": True, "2030-07-19": False, None: False})

# Note (27/10): free text for the whole payslip, from the "Note" column; empty = no note.
def _note_file(note, tag):
	wb_ = Workbook()
	ws_ = wb_.active
	ws_.append(list(payslip_admin.TEMPLATE_IDENT) + comp_cols)
	ws_.append([1, "2030-07", None, None, None, EMP_A, "", "", "", "", "", note] + vals(good))
	buf_ = io.BytesIO()
	wb_.save(buf_)
	return frappe.get_doc({"doctype": "File", "file_name": f"unit-note{tag}-{RUN}.xlsx", "is_private": 1, "content": buf_.getvalue()}).insert(ignore_permissions=True)


NOTE = "Điều chỉnh tháng 6: +150.000\nLiên hệ C&B nếu có thắc mắc"
f7 = _note_file(NOTE, 1)
r7 = payslip_admin.parse_upload(f7.file_url)["rows"][0]
check("note: read from the Note column", r7["note"], NOTE)
check("note: change shown in the overwrite warning", any(w.startswith("Note:") for w in r7["warnings"]), True)
payslip_admin.import_upload(f7.file_url, confirm_warnings=1)
ps_note = frappe.get_doc("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"})
check("note: saved on the payslip", ps_note.note, NOTE)
from customize_erpnext.api.employee_portal.payslip import build_payslip_payload
check("note: in the Mini App payload", build_payslip_payload(ps_note, False)["note"], NOTE)
f8 = _note_file("x" * (payslip_admin.NOTE_MAX + 1), 2)
check("note: too long = row error", bool(payslip_admin.parse_upload(f8.file_url)["rows"][0]["errors"]), True)
payslip_admin.discard_upload(f8.file_url)
check("note: template has a Note column after Group", list(payslip_admin.TEMPLATE_IDENT)[-2:], ["Group", "Note"])

# Download PDF (27/10): A4, one page per payslip; several -> ZIP of one PDF each.
import zipfile as _zip
from pypdf import PdfReader
from customize_erpnext.employee_portal import payslip_print

print_names = frappe.get_all("TIQN Payslip", {"period": "2030-07"}, pluck="name", order_by="employee")
files = payslip_print.build_files(print_names)
check("print: one PDF per payslip, one A4 page each",
	[(len(PdfReader(io.BytesIO(pdf)).pages), round(float(PdfReader(io.BytesIO(pdf)).pages[0].mediabox.height) / 72 * 25.4)) for _d, _n, pdf in files],
	[(1, 297)] * len(print_names))
pdf_a = next(pdf for d, _n, pdf in files if d.employee == EMP_A)
text_a = PdfReader(io.BytesIO(pdf_a)).pages[0].extract_text()
check("print: Withdrawn payslip has the DRAFT watermark", "DRAFT" in text_a, True)
check("print: note printed", "Liên hệ C&B" in text_a, True)
check("print: Vietnamese number format", (payslip_print._fmt_money(17214001), payslip_print._fmt_qty(1234.5), payslip_print._fmt_qty(26)),
	("17.214.001", "1.234,5", "26"))
_net_rows = [r for r in payslip_print.payslip_context(frappe.get_doc("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"}))["rows"] if r.is_net]
check("print: net pay line (Report Column Net) has VNĐ as its unit", [r.unit for r in _net_rows], ["VNĐ"])
check("print: file name {Period}_{section}_{group}_{employee}_{employee_name}.pdf",
	payslip_print.pdf_filename(frappe._dict(period="2030-07", section="HR-GA", group="IT", employee=EMP_A, employee_name="Nguyễn A/B")),
	f"2030-07_HR-GA_IT_{EMP_A}_Nguyễn A B.pdf")
check("print: empty parts left out", payslip_print.pdf_filename(frappe._dict(period="2030-07", section=None, group="", employee=EMP_A, employee_name="X")),
	f"2030-07_{EMP_A}_X.pdf")
payslip_print.download_pdf(frappe.as_json(print_names))
with _zip.ZipFile(io.BytesIO(frappe.response["filecontent"])) as zf:
	check("print: several -> ZIP with one PDF each", (frappe.response["filename"], len(zf.namelist())), (f"2030-07_{len(print_names)} payslips.zip", len(print_names)))
payslip_print.download_pdf(frappe.as_json(print_names[:1]))
check("print: one -> the PDF itself", frappe.response["filename"].endswith(".pdf"), True)
check("print: more than 200 refused", refused_fn(lambda: payslip_print.build_files([f"x{i}" for i in range(201)])), "refused")
frappe.set_user("Guest")
try:
	payslip_print.download_pdf(frappe.as_json(print_names[:1]))
	print_guest = "printed"
except frappe.PermissionError:
	print_guest = "refused"
frappe.set_user("Administrator")
frappe.clear_messages()
check("print: role required", print_guest, "refused")

# List Download PDF (27/10): ticked or a period, background job, ZIP (default) or one combined PDF, once.
_enqueued = []
_real_enqueue = frappe.enqueue
frappe.enqueue = lambda method, **kw: _enqueued.append(kw)
try:
	for _mode in ("zip", "combined"):
		_job = payslip_print.start_pdf_job(period="2030-07", mode=_mode)
		check(f"period pdf {_mode}: queued with every payslip", (_job["total"], _enqueued[-1]["mode"]), (len(print_names), _mode))
		check(f"period pdf {_mode}: one job at a time", refused_fn(lambda: payslip_print.start_pdf_job(period="2030-07")), "refused")
		payslip_print.build_pdf_job(_job["token"], _mode)
		_st = payslip_print.pdf_job_status(_job["token"])
		check(f"period pdf {_mode}: done", (_st["status"], _st["done"]), ("done", len(print_names)))
		frappe.set_user("Guest")
		try:
			payslip_print.pdf_job_status(_job["token"])
			_other = "seen"
		except (frappe.PermissionError, frappe.ValidationError):
			_other = "refused"
		frappe.set_user("Administrator")
		frappe.clear_messages()
		check(f"period pdf {_mode}: other users cannot see it", _other, "refused")
		payslip_print.download_pdf_job(_job["token"])
		_content = frappe.response["filecontent"]
		if _mode == "zip":
			with _zip.ZipFile(io.BytesIO(_content)) as zf:
				check("period pdf zip: one PDF per payslip", (frappe.response["filename"], len(zf.namelist())), (f"2030-07_{len(print_names)} payslips.zip", len(print_names)))
		else:
			check("period pdf combined: one PDF, one page per payslip", (frappe.response["filename"], len(PdfReader(io.BytesIO(_content)).pages)), (f"2030-07_{len(print_names)} payslips.pdf", len(print_names)))
		check(f"period pdf {_mode}: downloaded once", refused_fn(lambda: payslip_print.download_pdf_job(_job["token"])), "refused")
	_job = payslip_print.start_pdf_job(names=frappe.as_json(print_names), mode="combined")
	payslip_print.build_pdf_job(_job["token"], "combined")
	payslip_print.download_pdf_job(_job["token"])
	check("ticked pdf: combined PDF of the ticked payslips", (frappe.response["filename"], len(PdfReader(io.BytesIO(frappe.response["filecontent"])).pages)), (f"2030-07_{len(print_names)} payslips.pdf", len(print_names)))
	_job = payslip_print.start_pdf_job(names=frappe.as_json(print_names[:1]), mode="zip")
	payslip_print.build_pdf_job(_job["token"], "zip")
	payslip_print.download_pdf_job(_job["token"])
	check("ticked pdf: one payslip = its PDF even with ZIP picked", frappe.response["filename"].endswith(".pdf") and frappe.response["filecontent"][:4] == b"%PDF", True)
	check("pdf job: nothing ticked, no period refused", refused_fn(lambda: payslip_print.start_pdf_job()), "refused")
	check("pdf job options: periods latest first", payslip_print.pdf_job_options()["periods"][:1], [max(frappe.get_all("TIQN Payslip", pluck="period", distinct=True))])
	_emp_count = frappe.db.count("TIQN Payslip", {"employee": EMP_A})
	_emp_name = frappe.db.get_value("Employee", EMP_A, "employee_name")
	_job = payslip_print.start_pdf_job(employees=frappe.as_json([EMP_A]), mode="combined")
	check("employee pdf: no period = every payslip of the employee", _job["total"], _emp_count)
	payslip_print.build_pdf_job(_job["token"], "combined")
	payslip_print.download_pdf_job(_job["token"])
	check("employee pdf: file name {employee}_{name}_{n} payslips", frappe.response["filename"].endswith(f"{EMP_A}_{payslip_print._clean(_emp_name)}_{_emp_count} payslips.pdf"), True)
	_job = payslip_print.start_pdf_job(period="2030-07", employees=frappe.as_json([EMP_A]), mode="zip")
	payslip_print.build_pdf_job(_job["token"], "zip")
	payslip_print.download_pdf_job(_job["token"])
	check("employee pdf: period + one employee = that payslip",
		(_job["total"], frappe.response["filename"]),
		(1, payslip_print.pdf_filename(frappe.get_doc("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"}))))
	check("employee pdf: other employee, no payslip refused", refused_fn(lambda: payslip_print.start_pdf_job(employees=frappe.as_json(["zz-no-such-employee"]))), "refused")
	check("period pdf: Section filter", refused_fn(lambda: payslip_print.start_pdf_job(period="2030-07", sections=frappe.as_json(["zz-no-such-section"]))), "refused")
	_secs = sorted({s for s in frappe.get_all("TIQN Payslip", {"period": "2030-07"}, pluck="section") if s})
	if _secs:
		_job = payslip_print.start_pdf_job(period="2030-07", sections=frappe.as_json(_secs + ["zz-no-such-section"]), mode="zip")
		check("period pdf: several Sections", (_job["total"], frappe.cache.get_value(payslip_print.JOB_KEY.format(_job["token"]))["filename"]),
			(frappe.db.count("TIQN Payslip", {"period": "2030-07", "section": ("in", _secs)}), f"2030-07_{len(_secs) + 1} sections_{_job['total']} payslips.zip"))
		frappe.cache.delete_value([payslip_print.JOB_KEY.format(_job["token"]), payslip_print.USER_KEY.format("Administrator")])
	# Employees box: not Exclude Employees, not those who left before the period.
	from customize_erpnext.customize_erpnext.doctype.attendance_calculation_setting.attendance_calculation_setting import get_excluded_employee_ids
	_found = lambda txt, **kw: [r["value"] for r in payslip_print.search_pdf_employees(txt, **kw)]
	_excl = sorted(get_excluded_employee_ids())
	if _excl:
		check("pdf employees: Exclude Employees left out", _excl[0] in _found(_excl[0]), False)
	_left = frappe.db.sql("""SELECT name FROM tabEmployee WHERE relieving_date <= '2026-09-26' ORDER BY relieving_date DESC LIMIT 20""", pluck=True)
	_left = next((e for e in _left if e not in _excl), None)
	if _left:
		check("pdf employees: left before the period left out", _left in _found(_left, period="2026-10"), False)
		check("pdf employees: no period = still listed (old payslips)", _left in _found(_left), True)
	check("pdf job: Employees + Sections refused (one of the three)", refused_fn(lambda: payslip_print.start_pdf_job(period="2030-07", employees=frappe.as_json([EMP_A]), sections=frappe.as_json(["HR-GA"]))), "refused")
	check("pdf job: Sections + Group refused", refused_fn(lambda: payslip_print.start_pdf_job(period="2030-07", sections=frappe.as_json(["HR-GA"]), group="IT")), "refused")
	check("period pdf: unknown output refused", refused_fn(lambda: payslip_print.start_pdf_job(period="2030-07", mode="docx")), "refused")
finally:
	frappe.enqueue = _real_enqueue
check("publish takes re-imported withdrawn", payslip_admin.publish_period("2030-07")["published"], 2)
# ignore_no_copy=False = what the form's Duplicate does (frappe.model.copy_doc in JS).
_dup = frappe.copy_doc(frappe.get_doc("TIQN Payslip", {"period": "2030-07", "status": "Published"}), ignore_no_copy=False)
_dup.period = "2031-01"  # same employee + period = same name: HR changes one of them
_dup.insert()
check("duplicate of a published payslip saves as a new Draft, no publish history",
	(_dup.status, _dup.published_on, _dup.published_by, _dup.first_viewed_on, _dup.import_batch, str(_dup.period_from)),
	("Draft", None, None, None, None, "2030-12-26"))
ps_note.reload()
ps_note.note = "changed after publishing"
check("note: locked on a published payslip", refused_fn(lambda: ps_note.save(ignore_permissions=True)), "refused")
# Pre-filled template (23/10): pick a period -> one row per employee who worked in it.
from openpyxl import load_workbook

opts = payslip_admin.get_template_periods()
check("template periods: 13, default = last ended", (len(opts["periods"]), opts["default"]),
	(13, utils.shift_period(utils.period_of_date(today()), -1)))
payslip_admin.download_template(period="2026-09")
check("prefilled template filename has period + stamp", frappe.response["filename"].startswith("Payslip 2026-09 "), True)
tws = load_workbook(io.BytesIO(frappe.response["filecontent"])).active
trows = [r for r in tws.iter_rows(min_row=2, values_only=True) if r[5]]
listed = {r[5] for r in trows}
expected = payslip_admin.template_employees("2026-09")
check("prefilled: one row per employee of the period", len(trows), len(expected))
check("prefilled: No. = 1, 2, 3...", [r[0] for r in trows[:3]], [1, 2, 3][: len(trows)])
row_a = next((r for r in trows if r[5] == EMP_A), None)
emp_a = frappe.db.get_value("Employee", EMP_A, ["employee_name", "designation", "custom_section", "custom_group"], as_dict=True)
check("prefilled: employee columns filled", row_a and (row_a[1], row_a[6], row_a[7], row_a[9], row_a[10]),
	("2026-09", emp_a.employee_name, emp_a.designation, emp_a.custom_section, emp_a.custom_group))
check("prefilled: period dates, Pay Date empty", row_a and (str(getdate(row_a[2])), str(getdate(row_a[3])), row_a[4]), ("2026-08-26", "2026-09-25", None))
left_before = frappe.db.get_value("Employee", {"relieving_date": ["<=", "2026-08-26"]}, "name")
check("prefilled: left before the period not listed", left_before in listed, False)
left_during = frappe.get_all("Employee", filters={"relieving_date": ["between", ["2026-08-27", "2026-09-25"]]}, pluck="name")
check("prefilled: left during the period not listed (final settlement)", set(left_during) & listed, set())
on_maternity = frappe.get_all(
	"Employee Maternity", filters={"maternity_from_date": ["<=", "2026-08-26"], "maternity_to_date": [">=", "2026-09-25"]}, pluck="employee"
)
check("prefilled: maternity leave for the whole period not listed", set(on_maternity) & listed, set())
from customize_erpnext.customize_erpnext.doctype.attendance_calculation_setting.attendance_calculation_setting import (
	get_excluded_employee_ids,
)

excluded_ids = get_excluded_employee_ids()
check("prefilled: Attendance Calculation Setting > Exclude Employees not listed", excluded_ids & listed, set())
with mock.patch.object(payslip_admin, "get_excluded_employee_ids", return_value={EMP_A}):
	check("prefilled: exclusion applied (patched list)", EMP_A in {e.name for e in payslip_admin.template_employees("2026-09")}, False)
joined_during = frappe.db.get_value("Employee", {"date_of_joining": ["between", ["2026-08-27", "2026-09-25"]], "relieving_date": ["is", "not set"]}, ["name", "date_of_joining"], as_dict=True)
if joined_during:
	row_j = next((r for r in trows if r[5] == joined_during.name), None)
	check("prefilled: joined during the period -> Period From = joining", row_j and getdate(row_j[2]), getdate(joined_during.date_of_joining))

# Resigned template (26/10): a pay period + relieving date range inside it (empty = the whole
# period). Period To = relieving - 1, Pay Date = relieving + 14; relieving on the 26th left out.
r_period = next(
	(p for p in (utils.shift_period(utils.period_of_date(today()), -i) for i in range(12))
	 if payslip_admin.resigned_employees(*utils.period_bounds(p))),
	utils.shift_period(utils.period_of_date(today()), -1),
)
r_from, r_to = utils.period_bounds(r_period)
payslip_admin.download_template(period=r_period, resigned=1)
check("resigned template filename", frappe.response["filename"].startswith(f"Payslip {r_period} Resigned "), True)
rrows = [r for r in load_workbook(io.BytesIO(frappe.response["filecontent"])).active.iter_rows(min_row=2, values_only=True) if r[5]]
want = frappe.get_all("Employee", filters={"relieving_date": ["between", [r_from, r_to]]}, fields=["name", "relieving_date", "date_of_joining"])
want = [e for e in want if not e.date_of_joining or getdate(e.date_of_joining) < getdate(e.relieving_date)]
want = [e for e in want if e.name not in excluded_ids]
# Last day = end of a period (relieving on the 26th): whole period worked -> monthly template only.
want = [e for e in want if getdate(add_days(e.relieving_date, -1)) != utils.period_bounds(utils.period_of_date(add_days(e.relieving_date, -1)))[1]]
check("resigned: relieving on the 26th left out (on the monthly template)",
	(payslip_admin._left_at_period_end(getdate("2026-10-26")), payslip_admin._left_at_period_end(getdate("2026-10-25")),
	 payslip_admin._left_at_period_end(getdate("2026-10-27"))), (True, False, False))
check("resigned: only those who left in the period", sorted(r[5] for r in rrows), sorted(e.name for e in want))
check("resigned: every row in the chosen period", {r[1] for r in rrows} <= {r_period}, True)
if rrows:
	e0 = next(e for e in want if e.name == rrows[0][5])
	last_day = add_days(e0.relieving_date, -1)
	check("resigned: To = relieving - 1, Pay = relieving + 14",
		(rrows[0][1], getdate(rrows[0][3]), getdate(rrows[0][4])),
		(r_period, getdate(last_day), getdate(add_days(e0.relieving_date, 14))))
check("resigned: period required", refused(lambda: payslip_admin.download_template(resigned=1, from_date=r_from, to_date=r_to)), "refused")
check("resigned: range outside the period refused",
	refused(lambda: payslip_admin.download_template(period=r_period, resigned=1, from_date=add_days(r_from, -1), to_date=r_to)), "refused")
check("resigned: From after To refused",
	refused(lambda: payslip_admin.download_template(period=r_period, resigned=1, from_date=r_to, to_date=r_from)), "refused")
check("template periods carry their dates", {k for k in payslip_admin.get_template_periods()["periods"][0]} >= {"from_date", "to_date"}, True)
check("prefilled: bad period refused", refused(lambda: payslip_admin.download_template(period="2026-13")), "refused")
theader = [c.value for c in tws[1]]
c12 = next(l for l in SLINES if l.code == "1.2")
check("template: days column just before its amount", theader.index(payslip_admin.qty_header(c12)) + 1, theader.index(payslip_admin.amount_header(c12)))
check("template: days / hours headers short", (payslip_admin.qty_header(c12),
	payslip_admin.qty_header(next(l for l in SLINES if l.code == "2.1"))), ("1.2 Days", "2.1 Hours"))
check("template: starts with No.", theader[:2], ["No.", "Period"])
check("template: Note column right after Group", theader[10:12], ["Group", "Note"])
_m, qty_errors, _missing = payslip_admin._map_columns(
	["Kỳ lương", "MSNV", "3.2 Technical allowance - Days"], SLINES, STRUCT
)
# Language-independent: earlier portal calls leave frappe.local.lang = "vi".
check("upload: days column on a line without one is an error", (len(qty_errors), "3.2" in qty_errors[0] if qty_errors else None), (1, True))
_m, old_errors, _missing = payslip_admin._map_columns(["Kỳ lương", "MSNV", f"{payslip_admin.amount_header(c12)} - Unit/ĐVT", payslip_admin.amount_header(c12)], SLINES, STRUCT)
check("upload: older 'Unit/ĐVT' header still read", (old_errors, sorted(x[2] for x in _m["components"].values())), ([], [False, True]))
_m, name_errors, _missing = payslip_admin._map_columns(["Kỳ lương", "MSNV", payslip_admin.amount_header(c12)], SLINES, STRUCT)
check("upload: 'working days' in a name is not a days column", [x[2] for x in _m["components"].values()], [False])
_m, _e, _missing = payslip_admin._map_columns(
	["Kỳ lương", "MSNV", "1.1 Standard Working Days", "1.2 Days", "2.1 hour",
	 "2.2 OT working hours on weekend/Số giờ làm thêm ngày cuối tuần - Hours", "2.2 OT working hours on weekend/Số giờ làm thêm ngày cuối tuần"],
	SLINES,
	STRUCT,
)
check("upload: v5 short headers + older bilingual ones",
	(_e, [(x[0].code, x[2]) for x in _m["components"].values()]),
	([], [("1.1", False), ("1.2", True), ("2.1", True), ("2.2", True), ("2.2", False)]))
meta_ps = frappe.get_meta("TIQN Payslip")
check("desk: two tabs General / Detail", [df.label for df in meta_ps.fields if df.fieldtype == "Tab Break"], ["General", "Detail"])
ps_order = [df.fieldname for df in meta_ps.fields]
on_general = {f: ps_order.index(f) < ps_order.index("tab_detail") for f in ("net_pay", "standard_days", "actual_days", "ot_hours", "gross_pay", "lines_html")}
check("desk: net pay + report columns on General, the lines on Detail", on_general,
	{"net_pay": True, "standard_days": True, "actual_days": True, "ot_hours": True, "gross_pay": True, "lines_html": False})
check("desk: net_pay shown, sc_* hidden", (meta_ps.get_field("net_pay").hidden, all(df.hidden for df in meta_ps.fields if df.fieldname.startswith("sc_"))), (0, True))

imported = frappe.get_doc("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"})
imp = {r.code: r for r in pss.snapshot_of(imported)}
check("import writes the snapshot + report columns", (imp["1"].amount, imp["9"].amount, imp["1.2"].quantity, imported.net_pay, imported.structure),
	(good[0]["1"], good[0]["9"], good[1]["1.2"], good[0]["9"], STRUCT))
summary = next(p for p in payslip_admin.get_period_summary() if p["period"] == "2030-07")
check("period summary: viewed count of published", (summary["Published"], summary["viewed"]), (2, 1))

# Final settlement (24/10): left during the period -> To = relieving - 1, paid 14 days later,
# publishing it sends no notification.
leaver = next(
	(
		e
		for e in frappe.get_all(
			"Employee", filters={"relieving_date": ["is", "set"]}, fields=["name", "relieving_date", "date_of_joining"], order_by="relieving_date desc", limit=50
		)
		if e.date_of_joining
		and getdate(e.date_of_joining) < utils.period_bounds(utils.period_of_date(add_days(e.relieving_date, -1)))[0]
		and not frappe.db.exists("TIQN Payslip", {"period": utils.period_of_date(add_days(e.relieving_date, -1))})
	),
	None,
)
if leaver:
	fs_period = utils.period_of_date(add_days(leaver.relieving_date, -1))
	fs = frappe.get_doc({"doctype": "TIQN Payslip", "employee": leaver.name, "period": fs_period, "import_batch": f"UNIT-{RUN}"}).insert(ignore_permissions=True)
	check("final settlement: flagged + dates", (fs.final_settlement, getdate(fs.period_to), getdate(fs.pay_date), getdate(fs.period_from)),
		(1, getdate(add_days(leaver.relieving_date, -1)), getdate(add_days(leaver.relieving_date, 14)), utils.period_bounds(fs_period)[0]))
	res_fs = payslip_admin.publish_period(fs_period)
	check("final settlement: published without notification", (res_fs["published"], res_fs["final_settlement"], res_fs["notified"]), (1, 1, False))
	check("final settlement: no notification row", frappe.db.exists("TIQN Portal Notification", {"type": "PAYSLIP_PUBLISHED", "target_id": fs_period}), None)
	bad_dates = frappe.get_doc({"doctype": "TIQN Payslip", "employee": EMP_A, "period": "2030-10", "period_from": "2030-01-01"})
	check("period dates outside the period refused", refused(lambda: bad_dates.insert(ignore_permissions=True)), "refused")
check("monthly payslip: not a final settlement", frappe.db.get_value("TIQN Payslip", {"employee": EMP_A, "period": "2030-07"}, "final_settlement"), 0)

st = frappe.get_doc("Employee Portal Settings")
check("settings: logo field gone", st.meta.has_field("logo"), False)
check("settings: Day / Hour table gone (units from the field labels)", st.meta.has_field("payslip_qty_units"), False)
_l = {l.code: l for l in SLINES}
check("template: headers = Payslip Structure lines", theader[len(payslip_admin.TEMPLATE_IDENT):][:3],
	[payslip_admin.amount_header(_l["1"]), payslip_admin.amount_header(_l["1.1"]), payslip_admin.qty_header(_l["1.2"])])


# ---------------------------------------------------------------------------
print("\n=== M2: home / attendance / leave (read) ===")
from customize_erpnext.employee_portal import attendance_data as ad
from customize_erpnext.employee_portal import documents

for f in ("attendance", "overtime", "leave_view", "explanation_create"):
	frappe.db.set_single_value("Employee Portal Settings", f"feature_{f}", 1)
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")

r = call(api.get_home, token=TOKEN_A)
check("home: ok", r["ok"], True)
check("home: today card", bool(r["data"]["today"] and r["data"]["today"]["state"]), True)
check("home: annual leave", r["data"]["annual_leave"] is not None, True)
check("home: every contract key", sorted(r["data"]), sorted(["employee", "access", "today", "annual_leave", "attendance", "overtime", "latest_payslip", "notifications", "oa", "chat"]))

r = call(api.get_attendance_calendar, token=TOKEN_A, month="2026-09")
days = r["data"]["days"]
check("calendar: every day of the month", len(days), 30)
check("calendar: 02/09 is a holiday", days[1]["display"], "HOLIDAY")
check("calendar: holiday label translated", days[1]["label"], "Quốc khánh")
check("calendar: displays in the closed set", {d["display"] for d in days} <= {"PRESENT", "LATE_EARLY", "LEAVE", "HALF_LEAVE", "ABSENT", "HOLIDAY", "DAY_OFF", "NO_DATA", "TODAY", "FUTURE"}, True)
check("calendar: bad month", code(call(api.get_attendance_calendar, token=TOKEN_A, month="2026-13")), "VALIDATION")

r = call(api.get_attendance_day, token=TOKEN_A, date=add_days(today(), 3))
check("future day cannot be explained", r["data"]["explain_blocked"]["code"], "FUTURE_DATE")
r = call(api.get_attendance_day, token=TOKEN_A, date="2026-09-08")
check("day: punches carry direction", all(p["direction"] in ("IN", "OUT") for p in r["data"]["punches"]), True)
check("day: no approver -> blocked", r["data"]["explain_blocked"]["code"] if r["data"]["explain_blocked"] else None,
      None if frappe.db.get_value("Employee", EMP_A, "leave_approver") else "NO_APPROVER")

r = call(api.get_attendance_summary, token=TOKEN_A, period="2030-03")
# 2030-03 has a PUBLISHED payslip: totals still come from Attendance (D14 dropped 15/10).
check("summary: published period still from Attendance", (r["data"]["payable_days"]["source"], r["data"]["overtime"]["source"]), ("ERP", "ERP"))
check("summary: payable NOT taken from payslip", r["data"]["payable_days"]["value"] == qty_a["1.2"], False)
check("summary: cutoff = 25th", r["data"]["cutoff_date"], "2030-03-25")
check("summary: adjust until day before pay day", r["data"]["adjust_until"], "2030-04-04")
check("summary: future period provisional", r["data"]["provisional"], True)
check("provisional until adjust_until (04/10)", ad.period_summary(EMP_A, "2026-09", today="2026-10-04")["provisional"], True)
check("closed after adjust_until", ad.period_summary(EMP_A, "2026-09", today="2026-10-05")["provisional"], False)
check("published fake payslip does not close the period", ad.period_summary(EMP_A, "2026-07", today="2026-07-30")["provisional"], True)
s_att = ad.period_summary(EMP_A, "2026-09")
att_ot = frappe.db.sql("""SELECT SUM(custom_final_overtime_duration) FROM tabAttendance WHERE employee=%s AND docstatus<2
	AND attendance_date BETWEEN '2026-08-26' AND '2026-09-25'""", EMP_A)[0][0]
check("summary OT = sum of Attendance in period", s_att["ot_hours"], round(flt(att_ot), 2))

# OT registrations: insert raw rows (no hooks - the OT hooks recalculate attendance).
ot = frappe.new_doc("Overtime Registration")
ot.update({"name": f"UNIT-OT-{RUN}", "request_date": "2030-03-10", "workflow_state": "Draft", "docstatus": 0})
ot.db_insert()
for i, (emp, begin, end) in enumerate(((EMP_A, "17:00:00", "19:30:00"), (EMP_B, "17:00:00", "18:00:00"))):
	row = frappe.new_doc("Overtime Registration Detail")
	row.update({"name": f"UNIT-OTD-{RUN}-{i}", "parent": ot.name, "parenttype": "Overtime Registration", "parentfield": "ot_employees",
	            "employee": emp, "date": "2030-03-10", "begin_time": begin, "end_time": end, "reason": "unit", "idx": i + 1})
	row.db_insert()
r = call(api.get_overtime, token=TOKEN_A, period="2030-03")
mine = [i for i in r["data"]["items"] if i["registration_id"] == ot.name]
expected_status = "REGISTERED" if ad.include_draft_ot() else "DRAFT"
check("OT: only own rows", len(mine), 1)
check("OT: draft status follows include_draft_ot", (mine[0]["status"], mine[0]["hours"], mine[0]["id"]), (expected_status, 2.5, ot.name))
check("OT: draft hours counted only when REGISTERED", r["data"]["registered_approved_hours"] >= 2.5, expected_status == "REGISTERED")

# OT merged into the calendar + day detail (contract 1.1.3).
r = call(api.get_attendance_day, token=TOKEN_A, date="2030-03-10")
regs = r["data"]["overtime_registrations"]
check("day: OT registrations of the day", [(x["id"], x["status"], x["begin"], x["end"], x["hours"]) for x in regs], [(ot.name, expected_status, "17:00", "19:30", 2.5)])
check("day: OT status label translated", regs[0]["status_label"], "Đã đăng ký" if expected_status == "REGISTERED" else "Bản nháp")
check("day: no attendance -> approved OT null", r["data"]["approved_overtime_hours"], None)
att_day = frappe.db.sql("""SELECT attendance_date, custom_final_overtime_duration f, custom_approved_overtime_duration a FROM tabAttendance
	WHERE employee=%s AND docstatus<2 AND custom_final_overtime_duration>0 ORDER BY attendance_date DESC LIMIT 1""", EMP_A, as_dict=True)
if att_day:
	dd = att_day[0]
	r = call(api.get_attendance_day, token=TOKEN_A, date=str(dd.attendance_date))
	check("day: approved OT from Attendance", r["data"]["approved_overtime_hours"], round(flt(dd.a), 2))
	month = str(dd.attendance_date)[:7]
	r = call(api.get_attendance_calendar, token=TOKEN_A, month=month)
	day = next(x for x in r["data"]["days"] if x["date"] == str(dd.attendance_date))
	check("calendar: day OT = final OT", day["overtime_hours"], round(flt(dd.f), 2))
	sums = frappe.db.sql("""SELECT SUM(custom_final_overtime_duration), SUM(custom_approved_overtime_duration) FROM tabAttendance
		WHERE employee=%s AND docstatus<2 AND attendance_date LIKE %s""", (EMP_A, month + "-%"))[0]
	check("calendar: OT summary = month sums", (r["data"]["overtime_summary"]["actual_hours"], r["data"]["overtime_summary"]["approved_hours"]),
	      (round(flt(sums[0]), 2), round(flt(sums[1]), 2)))
	check("calendar: no-OT day -> null", all(x["overtime_hours"] is None or x["overtime_hours"] > 0 for x in r["data"]["days"]), True)
frappe.db.set_single_value("Employee Portal Settings", "feature_overtime", 0)
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
r = call(api.get_attendance_calendar, token=TOKEN_A, month="2030-03")
check("overtime feature off: no OT in calendar", r["data"]["overtime_summary"], None)
check("overtime feature off: no OT registrations on day", call(api.get_attendance_day, token=TOKEN_A, date="2030-03-10")["data"]["overtime_registrations"], [])
frappe.db.set_single_value("Employee Portal Settings", "feature_overtime", 1)
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
frappe.db.set_value("Overtime Registration", ot.name, {"workflow_state": "Rejected", "docstatus": 1}, update_modified=False)
r = call(api.get_overtime, token=TOKEN_A, period="2030-03")
check("OT: rejected registration hidden", any(i["registration_id"] == ot.name for i in r["data"]["items"]), False)

r = call(api.get_leave_balances, token=TOKEN_A)
check("balances: annual present", r["data"]["annual"]["leave_type"] in ad.annual_leave_types(), True)
check("balances: available = remaining - pending", r["data"]["annual"]["available"], round(r["data"]["annual"]["remaining"] - r["data"]["annual"]["pending"], 2))
r = call(api.get_leave_applications, token=TOKEN_A, year=2026)
check("leave list: counts add up", r["data"]["counts"]["ALL"], sum(v for k, v in r["data"]["counts"].items() if k != "ALL"))
check("leave list: bad status", code(call(api.get_leave_applications, token=TOKEN_A, status="x")), "VALIDATION")

other_la = frappe.db.get_value("Leave Application", {"employee": ["!=", EMP_A]}, "name")
check("leave detail of someone else -> NOT_FOUND", code(call(api.get_leave_application, token=TOKEN_A, id=other_la)), "NOT_FOUND")
own_la = frappe.db.get_value("Leave Application", {"employee": EMP_A}, "name")
if own_la:
	check("leave detail of own", call(api.get_leave_application, token=TOKEN_A, id=own_la)["data"]["id"], own_la)

other_ar = frappe.db.get_value("Attendance Request", {"employee": ["!=", EMP_A]}, "name")
if other_ar:
	check("explanation of someone else -> NOT_FOUND", code(call(api.get_attendance_request, token=TOKEN_A, id=other_ar)), "NOT_FOUND")

print("\n=== signed document links ===")
if own_la:
	fdoc = frappe.get_doc({"doctype": "File", "file_name": f"unit-{RUN}.txt", "is_private": 1, "content": b"hello",
	                       "attached_to_doctype": "Leave Application", "attached_to_name": own_la}).insert(ignore_permissions=True)
	sig, _exp = documents.sign(fdoc.name, EMP_A)
	frappe.local.response = frappe._dict()
	api.download_document(sig=sig)
	content = frappe.local.response.get("filecontent")
	check("download: own file streams", content.decode() if isinstance(content, bytes) else content, "hello")
	refused = []
	for label, bad in (("tampered", sig[:-2] + ("AA" if not sig.endswith("AA") else "BB")),
	                   ("signed for B", documents.sign(fdoc.name, EMP_B)[0]),
	                   ("expired", documents.sign(fdoc.name, EMP_A, ttl=-5)[0])):
		try:
			api.download_document(sig=bad)
			refused.append((label, "served"))
		except frappe.PermissionError:
			refused.append((label, "refused"))
	check("download: tampered / other employee / expired refused", [x[1] for x in refused], ["refused"] * 3)


# ---------------------------------------------------------------------------
print("\n=== link request (NOT_FOUND -> employee code -> HR approves) ===")
from customize_erpnext.employee_portal.doctype.tiqn_portal_link_request.tiqn_portal_link_request import normalise_name

check("normalise_name", normalise_name("  Nguyễn  Văn ĐẠT "), "nguyen van dat")
digits = str.maketrans("abcdef", "123456")
free_phone, free_phone2 = "0900001" + RUN[:3].translate(digits), "0900002" + RUN[:3].translate(digits)
check("lr: test phones match nobody", utils.find_employees_by_phone(free_phone) + utils.find_employees_by_phone(free_phone2), [])
missing_code = next(f"TIQN-{i:04d}" for i in range(9999, 0, -1) if not frappe.db.exists("Employee", f"TIQN-{i:04d}"))
name_a = frappe.db.get_value("Employee", EMP_A, "employee_name")
zid_lr = f"unit-{RUN}-lr"
r = auth_mod.complete_login(zid_lr, free_phone)
check("lr: NOT_FOUND, no request yet", (r["result"], r["link_request"]), ("NOT_FOUND", None))


AVATAR = "https://s120-ava-talk.zadn.vn/unit-test.jpg"
with mock.patch.object(zalo_graph, "_secret", return_value="s"), mock.patch.object(
	zalo_graph, "_get", return_value={"id": 7, "name": " N ", "picture": {"data": {"url": AVATAR}}}
):
	check("zalo: user info with picture", zalo_graph.get_user_info("at", "k"), {"id": "7", "name": "N", "picture": AVATAR})
with mock.patch.object(zalo_graph, "_secret", return_value="s"), mock.patch.object(
	zalo_graph, "_get", return_value={"id": 7, "picture": {"data": {"url": "javascript:alert(1)"}}}
):
	check("zalo: non-https picture dropped, no name", zalo_graph.get_user_info("at", "k"), {"id": "7", "name": None, "picture": None})


def link(zid=zid_lr, **kwargs):
	info = {"id": zid, "name": "Zalo LR", "picture": AVATAR}
	with mock.patch.object(zalo_graph, "get_user_info", return_value=info):
		return call(api.link_employee_request, access_token="at", **kwargs)


check("lr: no access token", code(call(api.link_employee_request, access_token="", employee_id=EMP_A, full_name="Ab")), "ZALO_TOKEN_INVALID")
check("lr: bad employee code", code(link(employee_id="0148", full_name="Ab")), "VALIDATION")
check("lr: name too short", code(link(employee_id=EMP_A, full_name=" A ")), "VALIDATION")
check("lr: Zalo account never signed in", code(link(zid=f"unit-{RUN}-nobody", employee_id=EMP_A, full_name="Ab")), "LOGIN_AGAIN")
zid_lnp = f"unit-{RUN}-lnp"
auth_mod.complete_login(zid_lnp, None, "STORED_PHONE")
check("lr: account without phone", code(link(zid=zid_lnp, employee_id=EMP_A, full_name="Ab")), "VALIDATION")
r_unknown = link(employee_id=missing_code, full_name=name_a)
check("lr: unknown employee code", code(r_unknown), "NOT_FOUND")
r_wrong = link(employee_id=EMP_A, full_name="Nguyen Van Khac Ten")
check("lr: wrong name -> the same NOT_FOUND as an unknown code (codes cannot be probed)",
	(code(r_wrong), r_wrong["error"]["message"], r_wrong["error"]["field"]), ("NOT_FOUND", r_unknown["error"]["message"], None))
check("lr: refused attempts create no record", frappe.db.exists("TIQN Portal Link Request", {"account": zid_lr}), None)
LR = "TIQN Portal Link Request"
check("lr: employee signing in from an active account -> EMPLOYEE_LINKED", code(link(employee_id=EMP_A, full_name=name_a)), "EMPLOYEE_LINKED")
check("lr: refused attempts create nothing", frappe.db.exists(LR, {"account": zid_lr}), None)
taken = set(frappe.get_all("TIQN Portal Account", filters={"employee": ["is", "set"]}, pluck="employee"))
EMP_LR = next(
	e for e in frappe.get_all("Employee", filters={"status": "Active", "relieving_date": ["is", "not set"]}, pluck="name", order_by="name desc", limit=200)
	if e not in taken and e not in (EMP_A, EMP_B)
)
name_lr = frappe.db.get_value("Employee", EMP_LR, "employee_name")
phone_lr_before = frappe.db.get_value("Employee", EMP_LR, "custom_cell_phone_zalo")
# A client-sent zalo_user_id is dropped: the request belongs to whoever Zalo says.
r = link(employee_id=EMP_LR.lower(), full_name=name_lr.upper(), zalo_user_id=zid_a)
check("lr: request sent", (r["ok"], r["data"]["status"], r["data"]["link_request"]["status"]), (True, "PENDING", "PENDING"))
lr = r["data"]["link_request"]
check("lr: counts in the answer", (lr["request_count"], lr["requests_today"], lr["max_per_day"], lr["remaining_today"], lr["can_resend"], bool(lr["requested_at"])),
	(1, 1, 2, 1, True, True))
check("lr: answer does not reveal the employee name", normalise_name(name_lr) in normalise_name(frappe.as_json(r)), False)
check("lr: client zalo_user_id ignored", frappe.db.exists(LR, {"account": zid_a}), None)
req = frappe.get_doc(LR, {"account": zid_lr})
check("lr: stored from server data", (req.employee, utils.phone_key(req.phone), req.zalo_name, req.status), (EMP_LR, utils.phone_key(free_phone), "Zalo LR", "Pending"))
check("lr: name match ignores accents/capitals", req.name_match, 1)
check("lr: Zalo avatar from Zalo, shown as the form image", (req.zalo_avatar, req.meta.image_field), (AVATAR, "zalo_avatar"))
check("lr: ERP counts", (req.request_count, req.requests_today, bool(req.last_requested_on)), (1, 1, True))
check("lr: ID of the request = the Zalo phone", req.name, free_phone)
check("lr: Employee editable on the form", frappe.get_meta(LR).get_field("employee").read_only, 0)
req.employee = EMP_B
req.save(ignore_permissions=True)
check("lr: HR corrects the employee while Pending -> checks redone", (req.employee, req.name_match, bool(req.warning)), (EMP_B, 0, True))
req.employee = EMP_LR
req.save(ignore_permissions=True)

name_plain = normalise_name(name_lr)  # typed without accents
r = link(employee_id=EMP_LR, full_name=name_plain)
check("lr: second request while pending overwrites", (r["ok"], frappe.db.count(LR, {"account": zid_lr})), (True, 1))
req.reload()
check("lr: overwritten content + counts", (req.full_name, req.name_match, req.request_count, req.requests_today), (name_plain, 1, 2, 2))
check("lr: previous content kept in the history", frappe.db.exists("Version", {"ref_doctype": LR, "docname": req.name}) is not None, True)
lr = auth_mod.complete_login(zid_lr, None, "STORED_PHONE")["link_request"]
check("lr: login shows counts, no resend left", (lr["status"], lr["employee_id"], lr["request_count"], lr["remaining_today"], lr["can_resend"]), ("PENDING", EMP_LR, 2, 0, False))
r = link(employee_id=EMP_LR, full_name=name_lr)
check("lr: 3rd request in a day -> REQUEST_LIMIT", (code(r), (r.get("error") or {}).get("details", {}).get("max_per_day")), ("REQUEST_LIMIT", 2))
check("lr: refused request changes nothing", (frappe.db.get_value(LR, req.name, "full_name"), frappe.db.get_value(LR, req.name, "request_count")), (name_plain, 2))
frappe.db.set_value(LR, req.name, "last_requested_on", add_days(now_datetime(), -1), update_modified=False)
lr = auth_mod.complete_login(zid_lr, None, "STORED_PHONE")["link_request"]
check("lr: next day the count starts again", (lr["requests_today"], lr["remaining_today"], lr["can_resend"]), (0, 2, True))
req.reload()
check("lr: ERP shows Requests Today = 0 next day", (req.run_method("onload"), req.requests_today)[1], 0)

req.reject("  Không đúng người ")
lr = auth_mod.complete_login(zid_lr, None, "STORED_PHONE")["link_request"]
check("lr: rejected shown with reason", (lr["status"], lr["reject_reason"]), ("REJECTED", "Không đúng người"))
check("lr: reviewer recorded", (req.reviewed_by, bool(req.reviewed_on)), ("Administrator", True))
check("lr: a rejected request cannot be approved", refused_fn(req.approve), "refused")

r = link(employee_id=EMP_LR, full_name=name_lr)
check("lr: can send again after reject", r["ok"], True)
req.reload()
check("lr: same record, back to Pending, review cleared", (req.status, req.reviewed_by, req.reject_reason, req.request_count, req.requests_today, req.name_match),
	("Pending", None, None, 3, 1, 1))

# Another ACTIVE account signing in as that employee blocks the app and Approve; disabling it frees the employee.
acc_a = frappe.db.get_value("TIQN Portal Account", zid_a, ["employee", "disabled"], as_dict=True)
frappe.db.set_value("TIQN Portal Account", zid_a, "employee", EMP_LR, update_modified=False)
req.set_checks()
check("lr: warns about the employee's other Zalo account", zid_a in (req.warning or ""), True)
check("lr: approve refused while another account signs in as the employee", refused_fn(req.approve), "refused")
check("lr: app refuses too", code(link(employee_id=EMP_LR, full_name=name_lr)), "EMPLOYEE_LINKED")
frappe.db.set_value("TIQN Portal Account", zid_a, "disabled", 1, update_modified=False)
from customize_erpnext.employee_portal.doctype.tiqn_portal_link_request.tiqn_portal_link_request import linked_accounts
check("lr: a disabled account does not count", linked_accounts(EMP_LR, exclude=zid_lr), [])
frappe.db.set_value("TIQN Portal Account", zid_a, {"employee": acc_a.employee, "disabled": acc_a.disabled}, update_modified=False)

# The Zalo account's phone changed: the same request moves to the new number.
free_phone3 = "0900003" + RUN[:3].translate(digits)
check("lr: third test phone matches nobody", utils.find_employees_by_phone(free_phone3), [])
frappe.db.set_value("TIQN Portal Account", zid_lr, "phone", free_phone3, update_modified=False)
old_name = req.name
r = link(employee_id=EMP_LR, full_name=name_lr)
check("lr: Zalo phone changed -> the request is renamed to the new phone",
	(r["ok"], frappe.db.exists(LR, old_name), frappe.db.get_value(LR, {"account": zid_lr}, "name"), frappe.db.get_value(LR, free_phone3, "request_count")),
	(True, None, free_phone3, 4))

req = frappe.get_doc(LR, free_phone3)
req.approve()
check("lr: approve writes the Zalo phone", utils.phone_key(frappe.db.get_value("Employee", EMP_LR, "custom_cell_phone_zalo")), utils.phone_key(free_phone3))
check("lr: approve links the Zalo account at once (no new sign-in needed)",
	frappe.db.get_value("TIQN Portal Account", zid_lr, ["employee", "disabled", "auto_disabled"]), (EMP_LR, 0, 0))
req.reload()
req.employee = EMP_B
check("lr: employee locked once reviewed", refused_fn(lambda: req.save(ignore_permissions=True)), "refused")
r = auth_mod.complete_login(zid_lr, None, "STORED_PHONE")
check("lr: next login ACTIVE as that employee", (r["result"], (r["employee"] or {}).get("employee_id"), r["link_request"]), ("ACTIVE", EMP_LR, None))
check("lr: phone now matches -> LOGIN_AGAIN", code(link(employee_id=EMP_LR, full_name="Ab")), "LOGIN_AGAIN")
dup = frappe.get_doc({"doctype": LR, "account": zid_lr, "employee": EMP_B, "full_name": "Ab", "phone": free_phone})
check("lr: one record per Zalo account (unique)", refused_fn(lambda: dup.insert(ignore_permissions=True)), "refused")

# Approving is refused when the Zalo phone already leads to another employee.
zid_blk = f"unit-{RUN}-blk"
auth_mod.complete_login(zid_blk, None, "STORED_PHONE")
blocked = frappe.get_doc({"doctype": LR, "account": zid_blk, "employee": EMP_LR, "full_name": "Ab", "phone": phone_of(EMP_B)})
blocked.insert(ignore_permissions=True)
check("lr: warns when the phone belongs to someone else", EMP_B in (blocked.warning or ""), True)
check("lr: approve refused when the phone belongs to someone else", refused_fn(blocked.approve), "refused")

zid_rl = f"unit-{RUN}-rl"
auth_mod.complete_login(zid_rl, free_phone2)
check("lr: 11th attempt in a day -> RATE_LIMITED (stops probing codes)",
	[code(link(zid=zid_rl, employee_id=missing_code, full_name="Ab")) for _ in range(11)], ["NOT_FOUND"] * 10 + ["RATE_LIMITED"])


# ---------------------------------------------------------------------------
print("\n=== Zalo App Review (sample data for Zalo's reviewers) ===")
from customize_erpnext.employee_portal import demo


def review_settings(mode, phones):
	frappe.db.set_single_value("Employee Portal Settings", {"review_mode": mode, "review_phones": phones})
	frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")


for f in ("payslip", "attendance", "overtime", "leave_view", "notifications"):
	frappe.db.set_single_value("Employee Portal Settings", f"feature_{f}", 1)
rv_phone = f"0988{int(RUN, 16) % 1000000:06d}"
rv_other = f"0987{int(RUN, 16) % 1000000:06d}"
assert not utils.find_employees_by_phone(rv_phone) and not utils.find_employees_by_phone(rv_other)
review_settings(1, f"0831111111\n{rv_phone}")
zid_rv = f"unit-{RUN}-rv"

check("review: every authenticated endpoint has a sample handler",
	sorted(n for n in api.__all__ if n not in ("get_config", "login", "link_employee_request", "dev_login", "logout", "download_document") and n not in demo.HANDLERS), [])
r = auth_mod.complete_login(zid_rv, "84" + utils.phone_key(rv_phone))
check("review: listed Zalo phone -> ACTIVE with the sample employee", (r["result"], r["employee"]["employee_id"]), ("ACTIVE", demo.EMPLOYEE_ID))
TOKEN_RV = r["token"]
check("review: account enabled, last_result REVIEW, no employee",
	frappe.db.get_value("TIQN Portal Account", zid_rv, ["disabled", "auto_disabled", "last_result", "employee"]), (0, 0, "REVIEW", None))
check("review: session REVIEW without employee",
	frappe.db.get_value("TIQN Portal Session", {"token_hash": hash_token(TOKEN_RV)}, ["login_method", "employee"]), ("REVIEW", None))
check("review: stored Zalo phone signs in again", auth_mod.complete_login(zid_rv, None, "STORED_PHONE")["result"], "ACTIVE")
check("review: phone not on the list -> NOT_FOUND", auth_mod.complete_login(f"unit-{RUN}-rvx", rv_other)["result"], "NOT_FOUND")
review_settings(1, f"{rv_phone}\n{phone_of(EMP_A)}")
r = auth_mod.complete_login(f"unit-{RUN}-rva", phone_of(EMP_A))
check("review: listed phone of a real employee signs in as the employee", (r["result"], r["employee"]["employee_id"]), ("ACTIVE", EMP_A))
tok_rva = r["token"]

# The real endpoints must never run for a reviewer: make their data readers explode.
boom = mock.Mock(side_effect=AssertionError("real data read in review mode"))
with mock.patch("customize_erpnext.api.employee_portal.payslip._published", boom), mock.patch(
	"customize_erpnext.employee_portal.attendance_data.attendance_rows", boom
), mock.patch("customize_erpnext.employee_portal.attendance_data.leave_balances", boom), mock.patch(
	"customize_erpnext.api.employee_portal.auth.employee_payload", boom
):
	rv = {n: call(getattr(api, n), TOKEN_RV) for n in (
		"me", "get_home", "get_notifications", "get_payslip_periods", "get_attendance_calendar", "get_attendance_summary",
		"get_overtime", "get_attendance_requests", "get_leave_balances", "get_leave_applications")}
	check("review: all screens answer", {n: code(x) for n, x in rv.items() if code(x)}, {})
	periods = [i["period"] for i in rv["get_payslip_periods"]["data"]["items"]]
	ps = call(api.get_payslip, TOKEN_RV, period=periods[0])
	check("review: sample payslip, net = paper formula", (code(ps), ps["data"]["employee"]["employee_id"], ps["data"]["net_pay"]), (None, demo.EMPLOYEE_ID, 8871250))
	check("review: sample payslip has the 9 numbered lines", [x["no"] for x in ps["data"]["rows"] if "." not in x["no"]], [str(i) for i in range(1, 10)])
	check("review: payslip marked seen (is_new off)",
		[i["is_new"] for i in call(api.get_payslip_periods, TOKEN_RV)["data"]["items"]], [False, True, True])
	check("review: 3 sample payslips = the last 3 ended periods, newest first", periods,
		[utils.shift_period(utils.period_of_date(today()), -i) for i in (1, 2, 3)])
	nets = [call(api.get_payslip, TOKEN_RV, period=p_)["data"]["net_pay"] for p_ in periods]
	check("review: the 3 sample payslips differ", len(set(nets)), 3)
	check("review: a payslip notification per sample period",
		sorted(i["target"]["id"] for i in call(api.get_notifications, TOKEN_RV)["data"]["items"] if i["type"] == "PAYSLIP_PUBLISHED"), sorted(periods))
	check("review: real period of an employee -> NOT_FOUND", code(call(api.get_payslip, TOKEN_RV, period="2030-07")), "NOT_FOUND")
	ar = rv["get_attendance_requests"]["data"]["items"][0]["id"]
	la = rv["get_leave_applications"]["data"]["items"][0]["id"]
	check("review: detail screens", [code(call(api.get_attendance_request, TOKEN_RV, id=ar)), code(call(api.get_leave_application, TOKEN_RV, id=la)),
		code(call(api.get_attendance_day, TOKEN_RV, date=add_days(today(), -2)))], [None, None, None])
	check("review: a real Attendance Request id -> NOT_FOUND",
		code(call(api.get_attendance_request, TOKEN_RV, id=frappe.db.get_value("Attendance Request", {}, "name") or "x")), "NOT_FOUND")
	check("review: mark all read", call(api.mark_notifications_read, TOKEN_RV, all=1)["data"]["unread"], 0)
	check("review: OA follow kept on the reviewer's own account",
		(call(api.set_oa_follow, TOKEN_RV, action="FOLLOWED")["data"]["followed"], frappe.db.get_value("TIQN Portal Account", zid_rv, "followed_oa")), (True, 1))
check("review: real data readers never called", boom.call_count, 0)
check("review: no unsaved sample payslip reached the DB", frappe.db.count("TIQN Payslip", {"employee": demo.EMPLOYEE_ID}), 0)

real = {n: call(getattr(api, n), tok_rva)["data"] for n in ("me", "get_home", "get_payslip_periods")}
for n, data in real.items():
	check(f"review: {n} has the keys of the real answer", sorted(rv[n]["data"]), sorted(data))
check("review: employee object has the real keys", sorted(rv["me"]["data"]["employee"]), sorted(real["me"]["employee"]))
try:
	demo.respond("cancel_leave_application", {})
	unknown = "answered"
except Exception as e:
	unknown = getattr(e, "code", repr(e))
check("review: endpoint without a sample handler -> NOT_FOUND", unknown, "NOT_FOUND")

frappe.db.set_value("Employee", EMP_B, "custom_cell_phone_zalo", rv_phone, update_modified=False)
x = call(api.get_home, TOKEN_RV)
check("review: phone now belongs to an employee -> session ends", (code(x), x.get("error", {}).get("details", {}).get("reason")), ("SESSION_INVALID", "REVOKED"))
frappe.db.set_value("Employee", EMP_B, "custom_cell_phone_zalo", None, update_modified=False)
tok2 = auth_mod.complete_login(zid_rv, None, "STORED_PHONE")["token"]
review_settings(1, "0831111111")
check("review: phone taken off the list -> session ends", code(call(api.get_home, tok2)), "SESSION_INVALID")
check("review: ... and login is NOT_FOUND again", auth_mod.complete_login(zid_rv, None, "STORED_PHONE")["result"], "NOT_FOUND")
review_settings(1, rv_phone)
tok3 = auth_mod.complete_login(zid_rv, None, "STORED_PHONE")["token"]
review_settings(0, rv_phone)
check("review: Review Mode off -> session ends", code(call(api.get_home, tok3)), "SESSION_INVALID")
check("review: Review Mode off -> NOT_FOUND", auth_mod.complete_login(zid_rv, None, "STORED_PHONE")["result"], "NOT_FOUND")
check("review: is_review_phone off", utils.is_review_phone(rv_phone), False)


# ---------------------------------------------------------------------------
print("\n=== payslip structure (lines per period, 28/10/2026) ===")
from customize_erpnext.employee_portal.doctype.tiqn_payslip_structure import tiqn_payslip_structure as pss_doc
from customize_erpnext.employee_portal.report.payslip_detail.payslip_detail import execute as detail_report


def _structure(valid_from, lines):
	return frappe.get_doc({"doctype": "TIQN Payslip Structure", "valid_from": valid_from, "lines": [dict(l) for l in lines]})


def _bad(change):
	lines_ = [dict(l) for l in SLINES]
	change(lines_)
	return refused_fn(lambda: _structure("2032-02", lines_).insert(ignore_permissions=True))


check("structure: No. used twice refused", _bad(lambda ls: ls.append(dict(ls[-1]))), "refused")
check("structure: No. not numbers refused", _bad(lambda ls: ls[0].update(code="1a")), "refused")
check("structure: Under Line that does not exist refused", _bad(lambda ls: ls[1].update(parent_code="99")), "refused")
check("structure: no Net line refused", _bad(lambda ls: [l.update(report_key=None) for l in ls if l["report_key"] == "Net"]), "refused")
check("structure: two Net lines refused", _bad(lambda ls: ls[0].update(report_key="Net")), "refused")
check("structure: Day line with a Unit/ĐVT column refused", _bad(lambda ls: next(l for l in ls if l["code"] == "1.1").update(quantity_unit="Day")), "refused")
check("structure: two Section Totals in one block refused", _bad(lambda ls: next(l for l in ls if l["code"] == "4").update(row_type="Section Total")), "refused")
check("structure: no formula is checked (any text allowed)", _bad(lambda ls: next(l for l in ls if l["code"] == "9").update(formula_display="(anything + 8%)")), "saved")
frappe.db.delete("TIQN Payslip Structure", {"valid_from": "2032-02"})
check("structure: may not take over periods that have payslips (2030-07)",
	refused_fn(lambda: _structure("2030-06", SLINES).insert(ignore_permissions=True)), "refused")

# 2032-01: + 3.11, - 3.10, 3.2 renamed, new formula text on 9.
new_lines = [dict(l) for l in SLINES if l.code != "3.10"]
new_lines.append(dict(code="3.11", label_en="Night Shift Allowance", label_vi="Phụ cấp ca đêm", unit="VND", parent_code="3.1", row_type="Line", section="Earnings"))
next(l for l in new_lines if l["code"] == "3.2").update(label_vi="Phụ cấp tay nghề")
next(l for l in new_lines if l["code"] == "9").update(formula_display="(j = f - g + h + i + k)")
s31 = _structure("2032-01", new_lines).insert(ignore_permissions=True)
check("structure: formula text kept without its brackets", next(l.formula_display for l in s31.lines if l.code == "9"), "j = f - g + h + i + k")
check("structure: rows sorted by No. on save", [l.code for l in s31.lines], sorted([l.code for l in s31.lines], key=code_key))
check("structure: periods resolve", (pss.structure_for_period("2031-12"), pss.structure_for_period("2032-01"), pss.structure_for_period("2032-06")),
	(STRUCT, s31.name, s31.name))
L31 = pss.structure_lines(s31.name)
check("structure: 3.11 under 3.1 after 3.9", [l.code for l in L31 if l.code.startswith("3.")][-2:], ["3.9", "3.11"])

frappe.local.response = frappe._dict()
payslip_admin.download_template(period="2032-01")
from openpyxl import load_workbook as _lw
h31 = [c.value for c in _lw(io.BytesIO(frappe.response["filecontent"])).active[1]]
check("template 2032-01: + 3.11, - 3.10", ("3.11 Night Shift Allowance" in h31, any(str(x).startswith("3.10 ") for x in h31)), (True, False))

cols31 = []
for l in L31:
	if l.quantity_unit:
		cols31.append((payslip_admin.qty_header(l), l.code, True))
	cols31.append((payslip_admin.amount_header(l), l.code, False))
a31, q31 = _fake_amounts(EMP_A, "2032-01")
a31 = dict(a31, **{"3.11": 400_000})


def _file(rows_, tag, columns=None):
	columns = columns or cols31
	wb_ = Workbook()
	ws_ = wb_.active
	ws_.append(list(payslip_admin.TEMPLATE_IDENT) + [h for h, _c, _q in columns])
	for period_, emp_, a_, q_ in rows_:
		ws_.append([1, period_, None, None, None, emp_, "", "", "", "", "", None] + [(q_ if is_q else a_).get(c_, 0) for _h, c_, is_q in columns])
	buf_ = io.BytesIO()
	wb_.save(buf_)
	return frappe.get_doc({"doctype": "File", "file_name": f"unit-pss{tag}-{RUN}.xlsx", "is_private": 1, "content": buf_.getvalue()}).insert(ignore_permissions=True)


fm = _file([("2032-01", EMP_A, a31, q31), ("2030-11", EMP_B, a31, q31)], "mix")
check("upload: periods of two structures in one file refused", bool(payslip_admin.parse_upload(fm.file_url)["header_errors"]), True)
payslip_admin.discard_upload(fm.file_url)
fo = _file([("2032-01", EMP_A, a31, q31)], "old", columns=cols)
check("upload: old columns (3.10) for 2032-01 refused", any("3.10" in e for e in payslip_admin.parse_upload(fo.file_url)["header_errors"]), True)
payslip_admin.discard_upload(fo.file_url)
absurd = []
for kw in ({"q": {"1.2": -1}}, {"q": {"1.2": 32}}, {"q": {"2.1": 24 * 31 + 1}}, {"a": {"5": 10**11}}):
	fa = _file([("2032-01", EMP_A, dict(a31, **kw.get("a", {})), dict(q31, **kw.get("q", {})))], f"abs{len(absurd)}")
	absurd.append(len(payslip_admin.parse_upload(fa.file_url)["rows"][0]["errors"]))
	payslip_admin.discard_upload(fa.file_url)
check("upload: impossible numbers are errors (negative / > period days / > 24h x days / >= 100 bn)", absurd, [1, 1, 1, 1])
fl = _file([("2032-01", EMP_A, a31, dict(q31, **{"7.4": 45}))], "leave")
check("upload: days outside Attendance (7.4 leave paid out) may exceed the period", payslip_admin.parse_upload(fl.file_url)["rows"][0]["errors"], [])
payslip_admin.discard_upload(fl.file_url)
def _clean_formula(text):
	row_ = frappe.get_doc({"doctype": "TIQN Payslip Structure Line", "code": "1", "formula_display": text})
	pss.clean_line(row_)
	return row_.formula_display


check("structure: one pair of brackets stripped, \"(a) + (b)\" kept", [_clean_formula(x) for x in ("(c=a+b)", "(a) + (b)", "8%")],
	["c=a+b", "(a) + (b)", "8%"])
wb_np = Workbook()
wb_np.active.append(["MSNV", "1 Basic Salary"])
wb_np.active.append([EMP_A, 1])
_b = io.BytesIO()
wb_np.save(_b)
f_np = frappe.get_doc({"doctype": "File", "file_name": f"unit-np-{RUN}.xlsx", "is_private": 1, "content": _b.getvalue()}).insert(ignore_permissions=True)
check("upload: no Period column = one clear error", len(payslip_admin.parse_upload(f_np.file_url)["header_errors"]), 1)
payslip_admin.discard_upload(f_np.file_url)
dcols_no = next(c for c in detail_report({"period": "2030-07"})[0] if c["fieldname"] == "no")
check("detail report: No. not summed in the total row", dcols_no["fieldtype"], "Data")
f31 = _file([("2032-01", EMP_A, a31, q31)], "ok")
d31 = payslip_admin.parse_upload(f31.file_url)
check("upload 2032-01: clean, structure 2032", (d31["header_errors"], d31["missing_components"], d31["rows"][0]["errors"], d31["structure"]), ([], [], [], s31.name))
payslip_admin.import_upload(f31.file_url, confirm_warnings=1)
payslip_admin.publish_period("2032-01")
p31 = call(api.get_payslip, token=TOKEN_A, period="2032-01", include_zero=1)["data"]
r31 = {x["no"]: x for x in p31["rows"]}
check("app 2032-01: 3.11 shown, 3.10 gone, 3.2 renamed, new formula", ("3.11" in r31, "3.10" in r31, r31["3.2"]["label"], r31["9"]["formula"]),
	(True, False, "Phụ cấp tay nghề", "j=f-g+h+i+k"))
check("app 2032-01: 3.11 amount", r31["3.11"]["amount"], 400_000)
check("app 2032-01: sections still built (older app builds)", [s_["key"] for s_ in p31["sections"]][:2], ["EARNINGS", "DEDUCTIONS"])
r07 = {x["no"]: x for x in call(api.get_payslip, token=TOKEN_A, period="2030-07", include_zero=1)["data"]["rows"]}
check("app 2030-07 unchanged by the new structure", ("3.10" in r07, "3.11" in r07, r07["3.2"]["label"]), (True, False, "Phụ cấp kỹ thuật"))
per31 = next(i for i in call(api.get_payslip_periods, token=TOKEN_A, year=2032)["data"]["items"] if i["period"] == "2032-01")
check("content_version: periods item = payslip", (bool(p31["content_version"]), per31["content_version"]), (True, p31["content_version"]))
check("content_version: home latest_payslip", call(api.get_home, token=TOKEN_A)["data"]["latest_payslip"]["content_version"], p31["content_version"])

# Locked once used: labels only, then Apply Labels to Payslips.
s31.reload()
s31.append("lines", dict(code="3.12", label_en="X", label_vi="X", unit="VND", row_type="Line", section="Earnings"))
check("structure in use: adding a line refused", refused_fn(lambda: s31.save(ignore_permissions=True)), "refused")
s31.reload()
next(l for l in s31.lines if l.code == "3.11").report_key = "Basic"
check("structure in use: changing a Report Column refused", refused_fn(lambda: s31.save(ignore_permissions=True)), "refused")
s31.reload()
next(l for l in s31.lines if l.code == "3.11").label_vi = "Phụ cấp làm đêm"
check("structure in use: changing a label allowed", refused_fn(lambda: s31.save(ignore_permissions=True)), "saved")
ch = pss_doc.get_label_changes(s31.name)
check("apply labels: change listed", [(c["code"], c["old"], c["new"], c["payslips"]) for c in ch["changes"]], [("3.11", "Phụ cấp ca đêm", "Phụ cấp làm đêm", 1)])
before31 = frappe.db.get_value("TIQN Payslip", {"employee": EMP_A, "period": "2032-01"}, ["published_on", "net_pay", "modified", "content_version"], as_dict=True)
check("apply labels: 1 payslip changed", pss_doc.apply_labels(s31.name)["changed"], 1)
after31 = frappe.db.get_value("TIQN Payslip", {"employee": EMP_A, "period": "2032-01"}, ["published_on", "net_pay", "modified", "content_version"], as_dict=True)
check("apply labels: published_on, net, modified unchanged; content_version changed",
	(after31.published_on == before31.published_on, after31.net_pay == before31.net_pay, after31.modified == before31.modified, after31.content_version != before31.content_version),
	(True, True, True, True))
r31b = {x["no"]: x for x in call(api.get_payslip, token=TOKEN_A, period="2032-01")["data"]["rows"]}
check("apply labels: the app shows the new label", r31b["3.11"]["label"], "Phụ cấp làm đêm")
check("apply labels: nothing left to apply", pss_doc.get_label_changes(s31.name)["changes"], [])
from customize_erpnext.employee_portal import payslip_print as _pp
ctx31 = _pp.payslip_context(frappe.get_doc("TIQN Payslip", {"employee": EMP_A, "period": "2032-01"}))
row311 = next(r for r in ctx31["rows"] if r.no == "3.11")
check("print 2032-01: English / Vietnamese of the snapshot", (row311.english, row311.vietnamese), ("Night Shift Allowance", "Phụ cấp làm đêm"))
check("print 2032-01: formula text", next(r.formula for r in ctx31["rows"] if r.no == "9"), "j=f-g+h+i+k")
check("structure in use cannot be deleted", refused_fn(lambda: frappe.delete_doc("TIQN Payslip Structure", s31.name)), "refused")

dcols, drows = detail_report({"period": "2032-01"})
dfields = [c["fieldname"] for c in dcols]
check("detail report: columns of the period's lines", ("l_3_11" in dfields, "l_3_10" in dfields, "q_1_2" in dfields, "q_1_1" in dfields), (True, False, True, True))
check("detail report: values", (drows[0]["l_3_11"], drows[0]["q_1_2"]), (400_000, q31["1.2"]))
check("detail report: 2030-07 keeps 3.10", "l_3_10" in [c["fieldname"] for c in detail_report({"period": "2030-07"})[0]], True)

# Review Mode: the sample payslip has its own lines, never the real structure.
check("demo: own lines, not a structure", {l[0] for l in demo.DEMO_LINES} >= {"1", "9"} and not hasattr(demo, "structure_lines"), True)

# Delete Payslips (28/10): every Draft AND Withdrawn payslip of the period, never a Published one.
payslip_admin.withdraw_period("2032-01")
row32 = next(p_ for p_ in payslip_admin.get_period_summary() if p_["period"] == "2032-01")
check("summary: withdrawn + opened by the employee", (row32["Withdrawn"], row32["withdrawn_viewed"]), (1, 1))
pub_0303 = frappe.db.count("TIQN Payslip", {"period": "2030-03", "status": "Published"})
check("delete payslips: Published never deleted", (payslip_admin.delete_period("2030-03")["deleted"], frappe.db.count("TIQN Payslip", {"period": "2030-03", "status": "Published"})), (0, pub_0303))
del32 = payslip_admin.delete_period("2032-01")
check("delete payslips: Withdrawn deleted too", (del32["deleted"], del32["withdrawn"], del32["viewed"], del32["queued"]), (1, 1, 1, False))
check("delete payslips: none left in the period", frappe.db.count("TIQN Payslip", {"period": "2032-01"}), 0)
check("delete payslips: kept in Deleted Document", bool(frappe.db.exists("Deleted Document", {"deleted_doctype": "TIQN Payslip", "deleted_name": f"PS-2032-01-{EMP_A}"})), True)
make_payslip(EMP_B, "2030-06", status="Draft")
check("delete payslips: Draft deleted", payslip_admin.delete_period("2030-06")["deleted"], 1)
check("delete payslips: bad period refused", refused_fn(lambda: payslip_admin.delete_period("2030-13")), "refused")
check("structure free again once its payslips are gone", refused_fn(lambda: frappe.delete_doc("TIQN Payslip Structure", s31.name)), "saved")


# ---------------------------------------------------------------------------
print("\n=== payslip period summary report ===")
from customize_erpnext.employee_portal.report.payslip_period_summary.payslip_period_summary import execute as period_report

cols_r, rows_r, _m, _c, cards = period_report({"period": "2030-07"})
# Cards by position: labels are translated (the portal calls leave frappe.local.lang = "vi").
check("report: published / viewed / not viewed", [c["value"] for c in cards[:3]], [2, 1, 1])
check("report: main figures as columns", [c["fieldname"] for c in cols_r if c["fieldname"].endswith(("_pay", "_deductions"))],
	["basic_pay", "overtime_pay", "gross_pay", "total_deductions", "net_pay"])
check("report: No. column 1..n", [r.no for r in rows_r], [1, 2])
check("report: filter not viewed", len(period_report({"period": "2030-07", "viewed": "Not Viewed"})[1]), 1)
check("report: bad period refused", refused_fn(lambda: period_report({"period": "2030-13"})), "refused")
rp = payslip_admin.get_report_periods()
check("report periods filter: periods with payslips, newest first, default = newest",
	([x["value"] for x in rp["periods"]] == sorted({x["value"] for x in rp["periods"]}, reverse=True), "2030-07" in [x["value"] for x in rp["periods"]], rp["default"] == rp["periods"][0]["value"]),
	(True, True, True))


# ---------------------------------------------------------------------------
frappe.db.commit = _REAL_COMMIT
frappe.flags.in_test = False
_REAL_ROLLBACK()
# The rollback does not reach the Redis copy of the Settings (features edited above) -
# drop it, or the Mini App would serve the test values.
frappe.clear_document_cache("Employee Portal Settings", "Employee Portal Settings")
left = frappe.db.count("TIQN Portal Account", {"name": ["like", f"unit-{RUN}%"]}) + frappe.db.count(
	"TIQN Payslip", {"import_batch": f"UNIT-{RUN}"}
)
check("rollback left no UNIT records", left, 0)
check("rollback left no link requests", frappe.db.count("TIQN Portal Link Request", {"account": ["like", f"unit-{RUN}%"]}), 0)
check("rollback restored the Zalo phone", frappe.db.get_value("Employee", EMP_LR, "custom_cell_phone_zalo"), phone_lr_before)
print(f"\n{ok} passed, {fail} failed")
