# -*- coding: utf-8 -*-
"""Kiểm tra chữ ký / cam kết / sửa trên Desk / Excel "Info for Sign" — KHÔNG ghi DB.

Chỉ sửa doc trong bộ nhớ, gọi hàm thuần, cuối cùng rollback. Không gọi đường nào có commit
(save_form_data chỉ được gọi ở nhánh throw trước khi lưu).

    bench --site erp.tiqn.local execute \
        customize_erpnext.api.self_update_info.test_self_update_sign.run
"""

import io
import json

import frappe

from customize_erpnext.api.self_update_info import self_update_info_api as api

DT = "Employee Self Update Info"


def _expect_throw(fn):
	try:
		fn()
	except frappe.ValidationError:
		frappe.clear_messages()
		return True
	return False


def _read_sheet():
	from openpyxl import load_workbook

	wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
	ws = wb.active
	rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6)]
	return wb, ws, rows


def run():
	results = []

	def check(name, ok, detail=""):
		results.append((bool(ok), name, detail))

	config = api._build_config()
	fields = {f["fieldname"]: f for s in config["sections"] for f in s["fields"]}

	# ---------- dialog field list ----------
	fl = api.get_info_for_sign_fields()
	check("dialog: có mục Ghi chú cuối", fl[-1]["fieldname"] == api.INFO_FOR_SIGN_REMARKS, fl[-1])
	check("dialog: đủ field config", len(fl) - 1 == len([f for f in fields if f not in ("employee", "name")]), len(fl))

	# ---------- Info for Sign ----------
	recs = frappe.get_all(DT, filters={"status": "Submitted"}, fields=["name", "data_json"], limit=3, order_by="name")
	names = [r.name for r in recs]

	api._download_info_for_sign(json.dumps(names), None, 0)
	wb, ws, rows = _read_sheet()
	check("sign: 1 sheet tên Info for Sign", wb.sheetnames == ["Info for Sign"], wb.sheetnames)
	check(
		"sign: tiêu đề cột tiếng Việt",
		[c.value for c in ws[5]] == ["STT", "Nhân viên", "Thông tin", "Nội dung", "Chữ ký"],
	)
	check("sign: cột Chữ ký luôn trống", all(r[4] in (None, "") for r in rows))
	check("sign: A4 dọc", ws.page_setup.orientation == "portrait")
	first_who = next(r[1] for r in rows if r[1])
	check("sign: ô Nhân viên = Mã NV xuống dòng Họ tên", "\n" in first_who and first_who.split("\n")[0].startswith(("TIQN", "TT")), first_who)
	stts = [r[0] for r in rows if r[0] is not None]
	emps = [r[1] for r in rows if r[1]]
	check("sign: STT theo nhân viên 1..N", stts == list(range(1, len(emps) + 1)), (stts, len(emps)))
	check("sign: STT đứng cùng dòng Mã NV", all((r[0] is None) == (not r[1]) for r in rows))
	check("sign: tên file", frappe.response["filename"] == "employee_info_for_sign.xlsx")
	merged = [str(m) for m in ws.merged_cells.ranges]
	check("sign: có gộp ô Nhân viên", any(m.startswith("B") for m in merged), merged[:5])
	b_ranges = sorted(m.replace("B", "") for m in merged if m.startswith("B"))
	a_ranges = sorted(m.replace("A", "") for m in merged if m.startswith("A") and ":E" not in m)
	e_ranges = sorted(m.replace("E", "") for m in merged if m.startswith("E"))
	check("sign: Chữ ký gộp ô đúng như Nhân viên", e_ranges == sorted(m.replace("B", "") for m in merged if m.startswith("B")))
	check("sign: STT gộp ô đúng như Mã NV", a_ranges == b_ranges, (a_ranges, b_ranges))
	sl = api._sign_labels(config)
	check("sign: nhãn không trùng", len(set(sl.values())) == len(sl))
	items = api._sign_items(config, None)
	kinds = [k for _l, _f, k in items]
	check("sign: địa chỉ gộp theo Section", kinds.count("address") == 2, [l for l, _f, k in items if k == "address"])
	check("sign: liên hệ khẩn cấp 1 dòng", kinds.count("emergency") == 1)
	labels_all = {r[2] for r in rows}
	check("sign: không còn dòng Tỉnh/Xã lẻ", not any(" - Tỉnh" in l or " - Xã" in l for l in labels_all), labels_all)
	F = lambda fn, w="Auto": {"fieldname": fn, "widget": w}
	check(
		"sign: địa chỉ = thôn, xã, tỉnh",
		api._sign_join("address", [(F("p", "Address Province"), "Tỉnh Quảng Ngãi"), (F("w", "Address Ward"), "Xã Ba Tơ"), (F("v"), "Thôn 3")])
		== "Thôn 3, Xã Ba Tơ, Tỉnh Quảng Ngãi",
	)
	check(
		"sign: khẩn cấp = Tên (Quan hệ) - SĐT",
		api._sign_join("emergency", [(F("person_to_be_contacted"), "Nguyễn A"), (F("relation"), "Mẹ"), (F("emergency_phone_number"), "0905")])
		== "Nguyễn A (Mẹ) - 0905",
	)
	check(
		"sign: khẩn cấp thiếu quan hệ",
		api._sign_join("emergency", [(F("person_to_be_contacted"), "Nguyễn A"), (F("relation"), ""), (F("emergency_phone_number"), "0905")])
		== "Nguyễn A - 0905",
	)

	# Nội dung = mới, mới trống → cũ (tính lại theo từng dòng gộp)
	r0 = frappe.get_doc(DT, names[0])
	saved = json.loads(r0.data_json or "{}")
	expected = {}
	for label, fs, kind in api._sign_items(config, None):
		parts = []
		for f in fs:
			new_s = str(saved.get(f["fieldname"]) or "").strip()
			old_s = "" if f.get("custom") else str(frappe.db.get_value("Employee", r0.employee, f["fieldname"]) or "").strip()
			parts.append((f, api._sign_value_vi(f, new_s or old_s)))
		v = api._sign_join(kind, parts)
		if v:
			expected[label] = v
	got, cur = {}, None
	for r in rows:
		if r[1]:
			cur = r[1].split("\n")[0]
		if cur == r0.employee:
			got[r[2]] = r[3]
	expected.pop("Ghi chú", None); got.pop("Ghi chú", None)
	check("sign: nội dung = mới, trống thì cũ", got == expected, f"{len(got)} dòng của {r0.employee}")

	# chọn bớt field
	pick = [f for f in fields if f not in ("employee", "name")][:2]
	api._download_info_for_sign(json.dumps(names), json.dumps(pick), 0)
	_, _, rows2 = _read_sheet()
	labels2 = {r[2] for r in rows2}
	allowed = {l for l, _f, _k in api._sign_items(config, set(pick))}
	check("sign: chỉ field đã chọn", labels2 <= allowed, labels2)

	# chỉ thay đổi ⊆ tất cả
	api._download_info_for_sign(json.dumps(names), None, 1)
	_, _, rows3 = _read_sheet()
	check("sign: chỉ thay đổi ≤ tất cả", len(rows3) <= len(rows), (len(rows3), len(rows)))

	# All Info không đổi
	api.download_excel(json.dumps(names))
	from openpyxl import load_workbook

	wb_all = load_workbook(io.BytesIO(frappe.response["filecontent"]))
	check("All Info: New Data + Old Data", wb_all.sheetnames == ["New Data", "Old Data"], wb_all.sheetnames)

	# Select dịch vi
	sel = next((f for f in fields.values() if f["fieldtype"] == "Select"), None)
	check("Select: hàm dịch chạy", sel is None or api._sign_value_vi(sel, "L") == "L")
	check("Date dd/mm/yyyy", api._sign_value_vi({"fieldtype": "Date"}, "2001-05-14") == "14/05/2001")

	# ---------- validate_desk_edit ----------
	def fresh():
		d = frappe.get_doc(DT, names[0])
		d.load_doc_before_save()
		return d, json.loads(d.data_json or "{}")

	text_f = next((f for f in fields.values() if api._hr_editor(f) == "text"), None)
	d, data = fresh()
	if text_f:
		data[text_f["fieldname"]] = "0912345678" if text_f.get("validation") == "Phone" else "Giá trị test"
		d.data_json = json.dumps(data, ensure_ascii=False)
		check("desk: sửa text hợp lệ → qua", not _expect_throw(lambda: api.validate_desk_edit(d)), text_f["fieldname"])

	if sel and api._hr_editor(sel) == "select":
		d, data = fresh()
		data[sel["fieldname"]] = "KHÔNG-CÓ"
		d.data_json = json.dumps(data, ensure_ascii=False)
		check("desk: select sai option → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))

	pairs = api._address_pairs(config)
	if pairs:
		prov, ward = pairs[0]
		d, data = fresh()
		data[prov] = "Thành phố Hà Nội"
		data[ward] = "Xã Ba Tơ"
		d.data_json = json.dumps(data, ensure_ascii=False)
		check("desk: xã không thuộc tỉnh → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))
		d, data = fresh()
		data[prov] = "Tỉnh Quảng Ngãi"
		data[ward] = "Xã Ba Tô"
		d.data_json = json.dumps(data, ensure_ascii=False)
		check("desk: tỉnh/xã hợp lệ → qua", not _expect_throw(lambda: api.validate_desk_edit(d)))

	date_f = next((f for f in fields.values() if api._hr_editor(f) == "date"), None)
	check("desk: field Date có editor date", date_f is not None, date_f and date_f["fieldname"])
	if date_f:
		d, data = fresh()
		data[date_f["fieldname"]] = "2020-01-15" if date_f.get("validation") != "Future" else "2099-01-15"
		d.data_json = json.dumps(data, ensure_ascii=False)
		check("desk: sửa ngày hợp lệ → qua", not _expect_throw(lambda: api.validate_desk_edit(d)))
		d, data = fresh()
		data[date_f["fieldname"]] = "ngày-sai"
		d.data_json = json.dumps(data, ensure_ascii=False)
		check("desk: ngày sai định dạng → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))
		if date_f.get("validation") == "Past":
			d, data = fresh()
			data[date_f["fieldname"]] = "2099-01-01"
			d.data_json = json.dumps(data, ensure_ascii=False)
			check("desk: ngày cấp ở tương lai → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))

	d, data = fresh()
	data["khong_ton_tai"] = "x"
	d.data_json = json.dumps(data)
	check("desk: field lạ → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))

	d, _ = fresh()
	d.signature = api._SIGNATURE_PREFIX + "AAAA"
	check("desk: sửa chữ ký → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))

	d, _ = fresh()
	d.status = "Reviewed"
	check("desk: chỉ đổi status → qua", not _expect_throw(lambda: api.validate_desk_edit(d)))

	synced = frappe.db.get_value(DT, {"status": "Synced"}, "name")
	if synced and text_f:
		d = frappe.get_doc(DT, synced)
		d.load_doc_before_save()
		data = json.loads(d.data_json or "{}")
		data[text_f["fieldname"]] = "x"
		d.data_json = json.dumps(data)
		check("desk: phiếu Synced → chặn", _expect_throw(lambda: api.validate_desk_edit(d)))

	# ---------- save_form_data: chữ ký sai định dạng (throw TRƯỚC khi lưu) ----------
	check(
		"portal: chữ ký sai định dạng → chặn",
		_expect_throw(lambda: api.save_form_data(r0.employee, "{}", signature="data:text/html,<script>")),
	)
	check(
		"portal: chữ ký quá lớn → chặn",
		_expect_throw(lambda: api.save_form_data(r0.employee, "{}", signature=api._SIGNATURE_PREFIX + "A" * 500_000)),
	)

	# ---------- phiếu PDF/PNG ----------
	d = frappe.get_doc(DT, names[0])
	d.signature = api._SIGNATURE_PREFIX + "iVBORw0KGgo="
	d.commitment_confirmed = 1
	d.signed_on = frappe.utils.now_datetime()
	html = api._build_submission_html(d, json.loads(d.data_json or "{}"), config)
	check("receipt: có câu cam kết + ảnh chữ ký", api.COMMITMENT_TEXT in html and "class='sig'" in html)
	check("receipt: đã ký → không có Thời điểm gửi", "Thời điểm gửi" not in html and "Ký lúc" in html)
	d2 = frappe.get_doc(DT, names[0])
	html2 = api._build_submission_html(d2, json.loads(d2.data_json or "{}"), config)
	check("receipt: chưa ký → giữ Thời điểm gửi", "Thời điểm gửi" in html2 or not d2.submitted_on)
	d.signature = "javascript:alert(1)"
	html = api._build_submission_html(d, json.loads(d.data_json or "{}"), config)
	check("receipt: chữ ký lạ không chèn ảnh", "class='sig'" not in html)

	check("hr pdf: tên file", api._receipt_base(d).startswith(d.employee))
	api.download_pdf_for_hr(names[0])
	check("hr pdf: ra file PDF", frappe.response["filecontent"][:4] == b"%PDF" and frappe.response["filename"].endswith(".pdf"),
		(frappe.response["filename"], len(frappe.response["filecontent"])))

	cfg = api.get_field_config()
	check("config: có commitment_text + purpose_text", cfg.get("commitment_text") and cfg.get("purpose_text"))

	frappe.db.rollback()
	failed = [r for r in results if not r[0]]
	for ok, name, detail in results:
		print(f"{'PASS' if ok else 'FAIL'}  {name}  {detail if detail != '' else ''}")
	print(f"\n{len(results) - len(failed)}/{len(results)} passed")
	return {"passed": len(results) - len(failed), "failed": len(failed)}
