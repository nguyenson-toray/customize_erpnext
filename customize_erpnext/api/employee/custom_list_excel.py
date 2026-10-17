"""Xuất Generate Custom Employee List ra Excel (.xlsx).

Client (employee_list.js → generate_custom_employee_list) đã tự giải phạm vi
nhân viên + thứ tự sắp xếp, gửi lên danh sách mã theo đúng thứ tự cùng cấu hình
cột. Ở đây chỉ dựng workbook: cùng cấu trúc với bảng HTML in ra (STT, tiêu đề
nhóm gộp ô, "||" = xuống dòng, giới tính viết tắt, cột Hình ảnh nhúng ảnh thật).

Để ở file riêng (không nhét vào employee_utils.py) vì thêm hàm mới vào module
worker đã import thì phải `bench restart` mới gọi được.

⚠ Đã thử XlsxWriter embed_image() (ảnh "Place in Cell") ngày 06/10/2026 → file
mở ra bị lỗi, user yêu cầu quay lại ảnh nổi (openpyxl add_image) như dưới đây.
"""

import io
import json
import os
import re

import frappe
from frappe import _
from frappe.utils import getdate, nowdate

# Mirror CUSTOM_LIST_STANDARD_FIELDS bên JS: chỉ cho đọc đúng các field này.
ALLOWED_FIELDS = {
	"name", "employee_name", "attendance_device_id", "date_of_birth", "gender",
	"designation", "department", "custom_section", "custom_group",
	"date_of_joining", "relieving_date", "company_email", "cell_number", "image",
}
DERIVED_SOURCE = {"id_no_prefix": "name"}
DATE_FIELDS = {"date_of_birth", "date_of_joining", "relieving_date"}
GENDER_ABBREVIATION = {"Male": "M", "Female": "F", "Other": "O", "Nam": "M", "Nữ": "F", "Khác": "O"}

# Ảnh 15mm × 18.33mm như bản PDF ≈ 57 × 69 px ở 96dpi; lưu gấp đôi cho nét khi in.
PHOTO_PX = (57, 69)
PHOTO_ROW_HEIGHT_PT = 56
AUTO_COL_WIDTH = 25


def _col_width_chars(width):
	"""'70px' / '70' → độ rộng cột Excel (đơn vị ký tự ≈ px/7). '%' và trống → mặc định."""
	m = re.match(r"^\s*(\d+(?:\.\d+)?)\s*(px)?\s*$", width or "")
	if not m:
		return AUTO_COL_WIDTH
	return max(4, round(float(m.group(1)) / 7, 1))


def _photo_path(url):
	"""Đường dẫn file ảnh trên đĩa, chỉ trong thư mục files của site."""
	if not url or url.startswith(("http://", "https://")):
		return None
	if url.startswith("/private/files/"):
		base, rel = frappe.get_site_path("private", "files"), url[len("/private/files/"):]
	elif url.startswith("/files/"):
		base, rel = frappe.get_site_path("public", "files"), url[len("/files/"):]
	else:
		return None
	base = os.path.realpath(base)
	path = os.path.realpath(os.path.join(base, rel))
	if not path.startswith(base + os.sep) or not os.path.isfile(path):
		return None
	return path


def _photo_image(url):
	from openpyxl.drawing.image import Image as XLImage
	from PIL import Image, ImageOps

	path = _photo_path(url)
	if not path:
		return None
	try:
		with Image.open(path) as img:
			img = ImageOps.exif_transpose(img).convert("RGB")
			img.thumbnail((PHOTO_PX[0] * 2, PHOTO_PX[1] * 2))
			buf = io.BytesIO()
			img.save(buf, format="JPEG", quality=85)
			ratio = min(PHOTO_PX[0] / img.width, PHOTO_PX[1] / img.height)
			disp = (round(img.width * ratio), round(img.height * ratio))
	except Exception:
		return None
	buf.seek(0)
	xl_img = XLImage(buf)
	xl_img.width, xl_img.height = disp
	return xl_img


def _cell_value(emp, col):
	fieldname = col.get("fieldname")
	if not fieldname:
		return None
	if fieldname == "id_no_prefix":
		m = re.search(r"(\d+)\s*$", emp.get("name") or "")
		return m.group(1) if m else emp.get("name")
	v = emp.get(DERIVED_SOURCE.get(fieldname, fieldname))
	if v in (None, ""):
		return None
	if fieldname == "gender":
		return GENDER_ABBREVIATION.get(v, v)
	if fieldname in DATE_FIELDS:
		return getdate(v)
	return v


@frappe.whitelist()
def export_custom_employee_list_excel(employees, columns, title=None, orientation="portrait"):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	employees = json.loads(employees) if isinstance(employees, str) else employees
	columns = json.loads(columns) if isinstance(columns, str) else columns
	if not employees:
		frappe.throw(_("No employee selected"))
	for col in columns:
		fn = col.get("fieldname")
		if fn and DERIVED_SOURCE.get(fn, fn) not in ALLOWED_FIELDS:
			frappe.throw(_("Field {0} is not allowed").format(fn))

	db_fields = {"name"} | {DERIVED_SOURCE.get(c["fieldname"], c["fieldname"]) for c in columns if c.get("fieldname")}
	# get_list = có kiểm quyền đọc Employee; giữ đúng thứ tự client gửi lên.
	rows = frappe.get_list(
		"Employee", filters={"name": ["in", employees]}, fields=list(db_fields), limit_page_length=0
	)
	by_name = {r.name: r for r in rows}
	data = [by_name[n] for n in employees if n in by_name]

	wb = Workbook()
	ws = wb.active
	ws.title = "Employee List"

	thin = Side(style="thin", color="000000")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	head_fill = PatternFill("solid", fgColor="F0F0F0")
	head_font = Font(bold=True)
	center = Alignment(horizontal="center", vertical="center", wrap_text=True)

	ncols = len(columns) + 1
	has_group = any(c.get("group") for c in columns)

	# ── Tiêu đề ──
	ws.cell(row=1, column=1, value=title or "").font = Font(bold=True, size=14)
	ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
	ws.row_dimensions[1].height = 28

	h1, h2 = 3, 4 if has_group else 3
	header_rows = [h1, h2] if has_group else [h1]

	def head(r, c, text):
		cell = ws.cell(row=r, column=c, value=(text or "").replace("||", "\n"))
		cell.font, cell.fill, cell.alignment = head_font, head_fill, center

	head(h1, 1, "STT")
	if has_group:
		ws.merge_cells(start_row=h1, start_column=1, end_row=h2, end_column=1)
	i = 0
	while i < len(columns):
		col, c = columns[i], i + 2
		if not col.get("group"):
			head(h1, c, col.get("label"))
			if has_group:
				ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
			i += 1
			continue
		j = i
		while j < len(columns) and columns[j].get("group") == col.get("group"):
			head(h2, j + 2, columns[j].get("label"))
			j += 1
		head(h1, c, col["group"])
		if j - i > 1:
			ws.merge_cells(start_row=h1, start_column=c, end_row=h1, end_column=j + 1)
		i = j
	for r in header_rows:
		for c in range(1, ncols + 1):
			ws.cell(row=r, column=c).border = border
			ws.cell(row=r, column=c).fill = head_fill

	# ── Dữ liệu ──
	has_image = any(c.get("is_image") or c.get("fieldname") == "image" for c in columns)
	start = h2 + 1
	for idx, emp in enumerate(data):
		r = start + idx
		ws.cell(row=r, column=1, value=idx + 1).alignment = center
		ws.cell(row=r, column=1).border = border
		for ci, col in enumerate(columns):
			c = ci + 2
			cell = ws.cell(row=r, column=c)
			cell.border = border
			if col.get("fieldname") == "image":
				cell.alignment = center
				img = _photo_image(emp.get("image"))
				if img:
					ws.add_image(img, cell.coordinate)
				continue
			v = _cell_value(emp, col)
			cell.value = v
			if col.get("fieldname") in DATE_FIELDS:
				cell.number_format = "DD/MM/YYYY"
			cell.alignment = Alignment(vertical="center", wrap_text=not col.get("nowrap"))
		if has_image:
			ws.row_dimensions[r].height = PHOTO_ROW_HEIGHT_PT

	# ── Độ rộng cột + in ──
	ws.column_dimensions["A"].width = 5
	for ci, col in enumerate(columns):
		w = _col_width_chars(col.get("width"))
		if col.get("fieldname") == "image":
			w = max(w, 9)
		ws.column_dimensions[get_column_letter(ci + 2)].width = w

	ws.freeze_panes = ws.cell(row=start, column=1)
	ws.print_title_rows = f"{h1}:{h2}"
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.page_setup.orientation = "landscape" if orientation == "landscape" else "portrait"
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_margins.left = ws.page_margins.right = 0.3

	buf = io.BytesIO()
	wb.save(buf)

	safe_title = re.sub(r"[^\w\-]+", "_", title or "Employee_List", flags=re.UNICODE).strip("_")[:60]
	frappe.response["filename"] = f"{safe_title or 'Employee_List'}_{nowdate()}.xlsx"
	frappe.response["filecontent"] = buf.getvalue()
	frappe.response["type"] = "binary"
